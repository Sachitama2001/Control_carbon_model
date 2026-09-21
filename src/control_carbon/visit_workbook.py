"""Inspectable Excel representation of the pinned VISIT TKY daily pathway.

Expressions form an acyclic graph. The same expression tree emits Excel
formulas and evaluates cached values; no Excel iterative calculation is used.
This is a source-order daily-rate representation, not an exact continuous
embedding of the native map. Missing data remain missing, never zero-filled.
"""
from __future__ import annotations

import ast
from dataclasses import dataclass, field
import hashlib
import json
import math
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

import numpy as np

SOURCE_COMMIT = "3285bd8e131a932e338b59892751648fd9edcc7b"
SOURCE_DEFAULT = Path("/mnt/d/ct/VISIT-matrix/visit_local")
MISSING = "未設定"
INVALID = "要確認:定義域"
FUNCTIONS = {
    "MIN": min, "MAX": max, "EXP": math.exp, "LN": math.log,
    "LOG10": math.log10, "SQRT": math.sqrt, "SIN": math.sin,
    "COS": math.cos, "ASIN": math.asin, "ACOS": math.acos,
    "ATAN": math.atan, "ABS": abs,
}


@dataclass
class Node:
    name: str
    label: str
    unit: str
    sheet: str
    expression: str | None = None
    value: float | None = None
    source: str = ""
    note: str = ""
    tree: ast.AST | None = field(default=None, repr=False)


@dataclass
class State:
    name: str
    label: str
    unit: str
    group: str
    base: str
    source: str
    initial: float | None = None


@dataclass
class Flux:
    name: str
    donor: str | None
    recipients: dict[str, str]
    note: str = ""


class Graph:
    def __init__(self):
        self.nodes: dict[str, Node] = {}

    def add(self, name, label, unit, sheet, expression=None, value=None,
            source="", note=""):
        if name in self.nodes:
            raise ValueError(f"duplicate node {name}")
        tree = ast.parse(expression, mode="eval").body if expression is not None else None
        if tree is not None:
            refs = {n.id for n in ast.walk(tree) if isinstance(n, ast.Name)}
            missing = refs - self.nodes.keys() - FUNCTIONS.keys() - {"IF", "AND", "OR"}
            if missing:
                raise ValueError(f"{name} references undefined/forward nodes {missing}")
        self.nodes[name] = Node(name, label, unit, sheet, expression, value, source, note, tree)
        return name

    def evaluate(self, overrides=None):
        values = {}
        overrides = overrides or {}
        for name in overrides:
            if name not in self.nodes or self.nodes[name].expression is not None:
                raise ValueError(f"not an editable input: {name}")

        def ev(t):
            if isinstance(t, ast.Constant):
                return t.value
            if isinstance(t, ast.Name):
                v = values[t.id]
                if v is None:
                    raise LookupError(t.id)
                if isinstance(v, str):
                    raise ArithmeticError(v)
                return v
            if isinstance(t, ast.UnaryOp):
                return -ev(t.operand) if isinstance(t.op, ast.USub) else ev(t.operand)
            if isinstance(t, ast.BinOp):
                a, b = ev(t.left), ev(t.right)
                return {ast.Add: lambda: a+b, ast.Sub: lambda: a-b,
                        ast.Mult: lambda: a*b, ast.Div: lambda: a/b,
                        ast.Pow: lambda: a**b}[type(t.op)]()
            if isinstance(t, ast.Compare):
                a, b = ev(t.left), ev(t.comparators[0])
                return {ast.Lt: lambda: a<b, ast.LtE: lambda: a<=b,
                        ast.Gt: lambda: a>b, ast.GtE: lambda: a>=b,
                        ast.Eq: lambda: a==b, ast.NotEq: lambda: a!=b}[type(t.ops[0])]()
            if isinstance(t, ast.Call):
                fn = t.func.id
                if fn == "IF":
                    return ev(t.args[1]) if ev(t.args[0]) else ev(t.args[2])
                if fn == "AND":
                    return all([ev(a) for a in t.args])
                if fn == "OR":
                    return any([ev(a) for a in t.args])
                return FUNCTIONS[fn](*[ev(a) for a in t.args])
            raise TypeError(ast.dump(t))

        for name, n in self.nodes.items():
            try:
                if name in overrides:
                    if n.expression is not None:
                        raise ValueError(f"cannot override derived node {name}")
                    v = overrides[name]
                else:
                    v = n.value if n.tree is None else ev(n.tree)
                if v is not None and (isinstance(v, complex) or not math.isfinite(v)):
                    raise ArithmeticError("nonfinite")
                values[name] = v
            except LookupError:
                values[name] = None
            except (ArithmeticError, ValueError, TypeError):
                values[name] = INVALID
        return values

    def excel(self, name):
        """Every reference guards blanks/text; IF branches remain lazy in Excel."""
        def emit(t):
            if isinstance(t, ast.Constant):
                return str(int(t.value)) if isinstance(t.value, bool) else repr(t.value)
            if isinstance(t, ast.Name):
                return f"IF(ISNUMBER({t.id}),{t.id},NA())"
            if isinstance(t, ast.UnaryOp):
                return f"(-{emit(t.operand)})" if isinstance(t.op, ast.USub) else emit(t.operand)
            if isinstance(t, ast.BinOp):
                op = {ast.Add: "+", ast.Sub: "-", ast.Mult: "*", ast.Div: "/", ast.Pow: "^"}[type(t.op)]
                return f"({emit(t.left)}{op}{emit(t.right)})"
            if isinstance(t, ast.Compare):
                op = {ast.Lt:"<", ast.LtE:"<=", ast.Gt:">", ast.GtE:">=", ast.Eq:"=", ast.NotEq:"<>"}[type(t.ops[0])]
                return f"({emit(t.left)}{op}{emit(t.comparators[0])})"
            if isinstance(t, ast.Call):
                return t.func.id+"("+",".join(emit(a) for a in t.args)+")"
            raise TypeError(ast.dump(t))
        # Keep the formula itself in the formula bar; a comment also holds the
        # compact symbolic expression and the source location.
        return '=IFERROR('+emit(self.nodes[name].tree)+',"未設定/定義域")'


class TKYModel:
    def __init__(self, source=SOURCE_DEFAULT, config=None):
        self.source = Path(source)
        self.config = Path(config or self.source/"INPUT/Config_TKY.xlsx")
        self.graph = Graph()
        self.states: list[State] = []
        self.fluxes: list[Flux] = []
        self.issues: list[tuple[str, str, str, str]] = []
        self.source_text = {p.name:p.read_text(errors="replace") for p in self.source.glob("*.c")}
        self.source_text.update({p.name:p.read_text(errors="replace") for p in self.source.glob("*.h")})
        self._parameters()
        self._inputs_states()
        self._environment_water()
        self._physiology_phenology()
        self._plants()
        self._soil()
        self._nitrogen()
        self._diagnostics()
        self._metadata()
        self._matrix_nodes()

    def src(self, file, function=None, needle=None):
        text = self.source_text[file]
        match = re.search(r'^[ \t]*(?:void|double|int|long|short)\s+'+re.escape(function)+r'\s*\(',text,re.M) if function else None
        start = match.start() if match else (-1 if function else 0)
        if start < 0:
            raise ValueError((file, function))
        if needle:
            start = text.find(needle, start)
            if start < 0:
                raise ValueError((file, function, needle))
        line = text.count("\n",0,start)+1
        return f"visit_local/{file}:{line}"+(f"::{function}" if function else "")

    def add(self, name, expr, label, unit, sheet, file, function=None, note="", needle=None):
        return self.graph.add(name,label,unit,sheet,expression=expr,
                              source=self.src(file,function,needle),note=note)

    def flux(self, name, donor=None, recipients=None, note=""):
        self.fluxes.append(Flux(name,donor,recipients or {},note))

    def issue(self, code, subject, reason, needed):
        self.issues.append((code,subject,reason,needed))

    def _parameters(self):
        import openpyxl
        w = openpyxl.load_workbook(self.config, data_only=True)
        self.config_book = w
        if w['Site']['C4'].value != 'TKY' or w['Site']['C5'].value != 4 or w['Site']['C9'].value != 1:
            raise ValueError('This workbook implementation selects TKY / DBF / soil texture 1 only.')
        self.config_hash = hashlib.sha256(self.config.read_bytes()).hexdigest()
        site = [("lat",6,"緯度","deg"),("lon",7,"経度","deg"),
                ("alt",8,"標高","m"),("texture",9,"土性コード","1=medium"),
                ("fc30",10,"上層保水容量","mm"),("fc",11,"深層保水容量（原典名称 total）","mm"),
                ("rootdepth",12,"根圏深さ","mm"),("hydcond",13,"飽和透水係数","m s-1"),
                ("density",14,"土壌容積密度","g cm-3"),("ph",15,"土壌pH","1"),
                ("sand",16,"砂割合","1"),("clay",17,"粘土割合","1")]
        for name,r,label,unit in site:
            self.graph.add("p_"+name,label,unit,"02_パラメータ",value=w["Site"].cell(r,3).value,
                           source=f"Config_TKY.xlsx!Site!C{r}")
        # The site reader overwrites the supplied upper-layer capacity with
        # Saxton (1986), then converts total capacity to the deep-layer bucket.
        # Use the values that actually reach hydrology, while retaining raw
        # Config values in notes and the copied Config sheet.
        sand=self.graph.nodes['p_sand'].value*100
        clay=self.graph.nodes['p_clay'].value*100
        a_sw=math.exp(-4.396-.0715*clay-4.488e-4*sand**2-4.285e-5*sand**2*clay)*100
        b_sw=-3.14-.00222*clay**2-3.484e-5*sand**2*clay
        fc30=(33/a_sw)**(1/b_sw)*300
        raw_fc30=self.graph.nodes['p_fc30'].value
        raw_total=self.graph.nodes['p_fc'].value
        self.graph.nodes['p_fc30'].value=fc30
        self.graph.nodes['p_fc30'].source=self.src('init_site.c','f_init_site','grid->fieldcap30 = grid->field_cap')
        self.graph.nodes['p_fc30'].note=f'Config入力{raw_fc30} mmはSaxton式で上書きされる。実行時有効値。'
        self.graph.nodes['p_fc'].value=raw_total-fc30
        self.graph.nodes['p_fc'].label='深層保水容量（実行時バケット）'
        self.graph.nodes['p_fc'].source=self.src('init_site.c','f_init_site','grid->fieldcap = grid->fieldcap - grid->fieldcap30')
        self.graph.nodes['p_fc'].note=f'Configの全層容量{raw_total} mmからSaxton上層{fc30:.12g} mmを差し引く。'
        keys = [
            ("albedo","アルベド","1"),("alloc_ass","葉配分係数","1"),
            ("alloc_abg","残りの地上部配分率","1"),("phenoltype","フェノロジー型","コード"),
            ("ctmp_lfdsp","展葉積算の閾値温度","degC"),("ctmp_lfshd","落葉積算の閾値温度","degC"),
            ("phototype","光合成型","3=C3;4=C4"),("sla","比葉面積","cm2 gDM-1; 半面積化は別係数"),
            ("eK0","消光係数","1"),("lue0","量子収率基準","molCO2 molPhoton-1"),
            ("pmax","最大光合成速度","umolCO2 m-2 s-1"),("topt0","最適温度基準","degC"),
            ("tmin","光合成下限温度","degC"),("tmax","光合成上限温度","degC"),
            ("gs_b0","気孔コンダクタンス切片","mmolH2O m-2 s-1"),
            ("gs_b1","Leuning係数","混合単位;原典のppm/umol/mmol規約"),
            ("gs_b2","VPD半飽和係数","hPa"),("km_nstl","土壌水分制限係数","1"),
            ("kmci","CO2半飽和係数","ppmv"),("cmpcd0","CO2補償点基準","ppmv"),
            ("rgf","葉成長呼吸係数","1"),("rgc","幹成長呼吸係数","1"),("rgr","根成長呼吸係数","1"),
            ("rmf0","葉維持呼吸基準","10^-3 day-1"),("rmc_s","辺材呼吸係数","10^-3 day-1"),
            ("rmr_s","細根呼吸係数","10^-3 day-1"),("rmc_h","心材呼吸係数","10^-3 day-1"),
            ("rmr_h","太根呼吸係数","10^-3 day-1"),
            ("qTf0","葉Q10基準","1"),("qTc0","幹Q10基準","1"),("qTr0","根Q10基準","1"),
            ("f_sz_s","幹サイズ係数","MgC ha-1"),("f_sz_r","根サイズ係数","MgC ha-1"),
            ("lf0","葉更新率","day-1"),("lc0","幹更新率","day-1"),("lr0","根更新率","day-1"),
            ("dcd","落葉期日落葉率","day-1"),("root_dpt_a","根分布係数a","m-1"),
            ("root_dpt_b","根分布係数b","m-1")]
        for p,col in [("t",5),("g",6),("v",7)]:
            for r,(key,label,unit) in enumerate(keys,5):
                from openpyxl.utils import get_column_letter
                cell=f"{get_column_letter(col)}{r}"
                raw=w["Parameter-daily"].cell(r,4).value
                self.graph.add(f"p_{p}_{key}",{"t":"樹木","g":"C3林床","v":"C4仮想群"}[p]+"："+label,
                               unit,"02_パラメータ",value=w["Parameter-daily"][cell].value,
                               source=f"Config_TKY.xlsx!Parameter-daily!{cell}",note=f"原ラベル {raw}; C4面積重みはTKYで0")
        skeys=[("albedo0","裸地アルベド","1"),("sr_lf","葉リター分解","10^-3 day-1"),
               ("sr_lc","幹リター分解","10^-3 day-1"),("sr_lr","根リター分解","10^-3 day-1"),
               ("sr_ha","活性腐植分解","10^-3 day-1"),("sr_hi","中間腐植分解","10^-3 day-1"),
               ("sr_hp","難分解腐植分解","10^-3 day-1"),("kml","リター水分係数","1"),
               ("kmh","腐植水分係数","1"),("kmsl","リター通気係数","1"),("kmsh","腐植通気係数","1"),
               ("f_co2_lf","葉リター呼吸分率","1"),("f_co2_lc","幹リター呼吸分率","1"),
               ("f_co2_lr","根リター呼吸分率","1"),("f_hm_a","活性腐植移行分率","1"),
               ("f_hm_i","中間腐植移行分率","1"),("f_hm_p","難分解腐植移行分率","1")]
        for r,(key,label,unit) in enumerate(skeys,45):
            self.graph.add("p_s_"+key,label,unit,"02_パラメータ",value=w["Parameter-daily"].cell(r,5).value,
                           source=f"Config_TKY.xlsx!Parameter-daily!E{r}")
        constants = [
            ("dt",1,"日次更新の幅","day","setting.h",None,"YSTEP"),
            ("pi",3.141592653,"原典の円周率","1","definition.h",None,"#define PI"),
            ("dTr",0.0174533,"度からラジアン","rad deg-1","definition.h",None,"#define dTr"),
            ("rTd",57.29577951,"ラジアンから度","deg rad-1","definition.h",None,"#define rTd"),
            ("dmTc",2.2,"炭素から乾物への原典係数","gDM gC-1","definition.h",None,"#define dmTc"),
            ("ZAT",273.15,"絶対温度オフセット","K","definition.h",None,"#define ZAT"),
            ("LHT",695,"原典潜熱","Wh kg-1","definition.h",None,"#define LHT"),
            ("cpwet",0.2813,"遮断・蒸散の空気比熱","Wh kg-1 K-1","hydro_flows.c","pm_incep","cp ="),
            ("cpsoil",1014,"土壌蒸発の上書き比熱","単位要確認","hydro_flows.c","pm_evap","cp = 1014.0"),
            ("eta",0.0000224,"コンダクタンス単位換算","m s-1 per mmol m-2 s-1","hydro_flows.c","pm_incep","eta ="),
            ("comp1",0.407,"TKY Quercus構成率","1","location_init.c","f_loct_init","loct->comp_over1 = 0.407"),
            ("comp2",0.593,"TKY Betula構成率","1","location_init.c","f_loct_init","loct->comp_over2 = 0.593"),
            ("weight_t",1,"樹木面積重み","1","daily_scheme.c","daily_scheme","flux->gpp ="),
            ("weight_g",1,"TKY C3林床重み","1","location_proc.c","f_loct_proc","loct->funder_c3 ="),
            ("weight_v",0,"TKY C4面積重み","1","location_proc.c","f_loct_proc","loct->funder_c4 = 0.0"),
        ]
        for key,val,label,unit,file,fn,needle in constants:
            self.graph.add("c_"+key,label,unit,"03_ソース定数",value=val,source=self.src(file,fn,needle))
        self.issue("SRC01","水：baseflow二重控除","hydro_balance.cはbaseflowをswwから控除し、ro2へ加えて再控除する。原典どおり分けて表示。","意図または修正版の指定")
        self.issue("SRC02","水：PM土壌蒸発のcp=1014","遮断・蒸散は0.2813、蒸発のみ1014へ上書き。原典を保持。","比熱・時間単位の正式な解釈")
        self.issue("SRC03","N：単位の不整合","n_budgetはgN/haの状態に、mg/m2のNGAS出力を換算不足のまま控除。NH3質量とN質量、biofix年率も混在。","原典を修正するかの判断・正しい単位規約")
        self.issue("SRC04","C：展葉費用の二重計上候補","NSCから展葉時に10%を引き、EPP>0ならrfgへ同額を加え葉から再控除。原典を保持。","意図または修正版")
        self.issue("SRC05","C4面積0でも窒素状態が更新","Nの回収や落葉はn_budgetで面積重みを再適用しない。C4は炭素総量に算入しないが内部群として保持。","C4を完全停止する際の正式な設定")
        self.issue("SRC06","日次版と30分版","FLUX_SCHEME=0を今回の簡単な実験設定として選択。30分FvCB、熱収支、Newton解法は非選択一覧へ。","30分版も必要な場合は別構成で展開")

    def _inputs_states(self):
        inputs=[("doy","年内日（0始まり）","day","radiation.c"),
                ("ta","2m気温","degC","location_proc.c"),("ts","地表温度","degC","location_proc.c"),
                ("tl","10cm地温","degC","location_proc.c"),("th","深層地温","degC","location_proc.c"),
                ("rain","日降水","mm day-1","hydro_balance.c"),("cloud","雲量","0..1","radiation.c"),
                ("wind","10m風速","m s-1","hydro_flows.c"),("vp","実水蒸気圧","hPa","location_proc.c"),
                ("vpd","飽差","hPa","location_proc.c"),("co2","大気CO2","ppmv","atm_co2.c"),
                ("ch4","大気CH4","ppmv","ch4_oxy.c"),
                ("depo_no3","評価日のNO3沈着","gN ha-1 day-1","location_proc.c"),
                ("depo_nh4","評価日のNH4沈着","gN ha-1 day-1","location_proc.c"),
                ("tsoil_mean","年平均10cm地温（NGAS2）","degC","n2o_emit.c"),
                ("ngas2_aa","NGAS2の未初期化ローカル変数aa","不明","n2o_emit.c")]
        for key,label,unit,file in inputs:
            self.graph.add("z_"+key,label,unit,"04_気象入力",source=self.src(file),note="値未提供。0を仮定しない。")
        for p in ("t","g","v"):
            self.graph.add(f"h_{p}_gc","前日から保持する気孔コンダクタンス "+p,
                           "mmolH2O m-2 s-1","05_履歴入力",source=self.src("location_proc.c","f_loct_proc","f_hydrology("),
                           note="hydrologyは当日ecophysiologyより先に実行されるため必要")
            for key,label in [("gdd","展葉積算温度"),("cdd","落葉積算温度"),("flush","展葉開始後日数"),("shed","落葉開始後日数")]:
                self.graph.add(f"h_{p}_{key}",p+":"+label,"degC day" if key in ("gdd","cdd") else "day",
                               "05_履歴入力",source=self.src("phenology.c","f_growth_period"))
        for p,label in [("t","樹木"),("g","C3林床"),("v","C4仮想群")]:
            for organ,jp,base in [("fol","葉",f"p_{p}_lf0"),("stm","幹",f"p_{p}_lc0"),
                                  ("rot","根",f"p_{p}_lr0"),("nsc","非構造性炭素","0.03")]:
                self._state(f"x_{p}_{organ}",label+jp,"MgC ha-1","C",base,"initi_mass.c","initTree" if p=="t" else "initFloor",.001 if organ=="nsc" else .01)
        for pool,base in [("tf","lf"),("tc","lc"),("tr","lr"),("gf","lf"),("gc","lc"),("gr","lr"),("ha","ha"),("hi","hi"),("hp","hp")]:
            self._state("x_s_"+pool,"土壌炭素 "+pool,"MgC ha-1","C",f"p_s_sr_{base}/1000","initi_mass.c","initSoil",.01)
        for p,label in [("t","樹木"),("g","C3林床"),("v","C4仮想群")]:
            for organ,jp in [("can","葉群窒素"),("str","貯蔵窒素")]:
                self._state(f"x_{p}_n{organ}",label+jp,"gN ha-1（収支式準拠）","N","1","n_budget.c","n_budget",.001)
        for pool,label in [("mic","微生物"),("lit","リター"),("hum","腐植"),("no3","硝酸態"),("nh4","アンモニア態")]:
            self._state("x_n_"+pool,label+"窒素","gN ha-1（収支式準拠）","N","1","n_budget.c","n_budget",.001)
        for key,label in [("snow","積雪水当量"),("sw","上層土壌水"),("dw","深層土壌水")]:
            self._state("x_w_"+key,label,"mm","W","1","hydro_balance.c","f_hydrology")
        self._state("x_doc","溶存有機炭素濃度（補助状態）","mgC L-1","D","1","doc.c","f_doc_boyer",0)
        self._state("x_casa","CASA水分記憶（補助状態）","m3 m-3","D","1","hydro_flows.c","f_casa_moisture")
        self.issue("INPUT01","評価日・気象・大気組成","Config_TKY.xlsxに時系列の実値はない。気圧と放射は入力から計算する。","04_気象入力の空欄を埋める。vpとVPDを供給。")
        self.issue("INPUT02","現在の状態ベクトル","Configは現在炭素・N・水・DOCを定めない。初期化コード値は参考列のみで採用しない。","評価日のrestart/観測/任意に選んだ状態を10_Xへ入力")
        self.issue("INPUT03","前日の生理量・フェノロジー記憶","気孔コンダクタンス、GDD/CDD、展葉・落葉後日数が必要。","05_履歴入力へ前日値を供給")
        self.issue("INPUT04","窒素沈着の現用データ","f_n_depositは呼出しがコメント化。現用はdepo_*_model配列。TKY=28.7を勝手に代入しない。","評価日のdepo_no3/depo_nh4（gN/ha/day）")
        self.issue("UNDEF01","NGAS2 aa","n2o_emit.c::f_n2o_emit_ngas_2のjjj=(a2-th_a/aa+ss)でaa未初期化。式を保持し入力は空欄。N本体はNGAS1出力を使用。","変数aaの意図または修正版ソース")
        self.issue("UNDEF02","Curry水分分岐","0.2>=ps && ps<=100の条件では0.2<ps<=100が未定義。代替値は入れない。","該当分岐の正式な訂正")

    def _state(self,name,label,unit,group,base,file,fn,initial=None):
        source=self.src(file,fn)
        self.states.append(State(name,label,unit,group,base,source,initial))
        self.graph.add(name,label,unit,"10_X",source=source,note="現在値は未設定。初期化参考値を現在の森林値と混同しない。")

    def _environment_water(self):
        a=lambda name,expr,label,unit="1",file="radiation.c",fn=None: self.add(name,expr,label,unit,"20_気象放射",file,fn)
        a("e_ge","2*c_pi*z_doy/365","太陽年周位相")
        a("e_dec","0.006918-0.399912*COS(e_ge)+0.070257*SIN(e_ge)-0.006758*COS(2*e_ge)+0.000907*SIN(2*e_ge)-0.002697*COS(3*e_ge)+0.00148*SIN(3*e_ge)","太陽赤緯","rad","radiation.c","f_soldec")
        a("e_dl","2*ACOS(MAX(-1,MIN(1,-SIN(p_lat*c_dTr)*SIN(e_dec)/(COS(p_lat*c_dTr)*COS(e_dec)))))*c_rTd/15","日長","h","radiation.c","f_daylen")
        a("e_sinnoon","MAX(-1,MIN(1,SIN(p_lat*c_dTr)*SIN(e_dec)+COS(p_lat*c_dTr)*COS(e_dec)))","正午太陽高度sin")
        a("e_dtc","1.00011+0.034221*COS(e_ge)+0.00128*SIN(e_ge)+0.000719*COS(2*e_ge)+0.000077*SIN(2*e_ge)","地球太陽距離補正")
        a("e_black","MAX(0,MIN(1,0.803-0.34*z_cloud-0.458*z_cloud**2))","雲量から短波透過率")
        for h in range(48):
            a(f"e_sw{h:02d}",f"e_black*MAX(0,1367*e_dtc*(SIN(p_lat*c_dTr)*SIN(e_dec)+COS(p_lat*c_dTr)*COS(e_dec)*COS({-180+h*7.5}*c_dTr)))",f"{h/2:g}時 短波放射","W m-2")
        a("e_swmean","("+"+".join(f"e_sw{h:02d}" for h in range(48))+")/48","日平均短波（48点平均）","W m-2")
        a("e_hd","e_sw24*(0.958-0.982*e_black)","正午散乱短波","W m-2")
        a("e_parb","0.43*(e_sw24-e_hd)","直達PAR","W m-2")
        a("e_pard","0.57*e_hd","散乱PAR","W m-2")
        a("e_ppfdc","e_parb*4.6+e_pard*4.2","McCree正午PPFD（樹冠GPP用）","umolPhoton m-2 s-1")
        a("e_diff","MAX(0.01,0.958-0.982*e_black)","散乱割合")
        a("e_ppfd","e_parb*(4.576-0.033144*e_diff)+e_pard*MAX(4.2,4.5886*e_diff/(0.010773+e_diff))","Dye正午PPFD（単葉反復用）","umolPhoton m-2 s-1")
        a("e_pressure","1013.25*EXP(-0.028964*9.8*MAX(0,p_alt)/(8.3144*(z_ta+c_ZAT)))","気圧","hPa","location_proc.c","f_loct_proc")
        a("e_vps","6.1078*10**IF(z_ta>0,7.5*z_ta/(237.3+z_ta),9.5*z_ta/(265.3+z_ta))","飽和水蒸気圧","hPa","hydro_flows.c","f_vap_pre_sat")
        a("e_slope","e_vps*IF(z_ta>0,2500-2.4*z_ta,2834)/(0.4615*(c_ZAT+z_ta)**2)","飽和蒸気圧勾配","hPa K-1","hydro_flows.c","f_slope_vps")
        a("e_density","1.293*c_ZAT/(z_ta+c_ZAT)*e_pressure/1013.25*(1-0.378*z_vp/e_pressure)","空気密度","kg m-3","hydro_flows.c","f_airdens")
        a("e_ra","MIN(59.5,MAX(0.1,LN(10)**2/(0.41**2*MAX(0.1,z_wind))))","空気力学抵抗","s m-1","hydro_flows.c","f_r_aero")
        # TKY SLA changes before daily phenology; use previous-day counters.
        for tag,start,slope,base,late in [("one",50,2,78,108),("two",60,.12,61.8,76.68)]:
            mid=.24 if tag=="one" else .12
            end=88 if tag=="one" else 56.68
            expr=f"IF(AND(h_t_flush<=0,z_doy<210),{start},IF(AND(h_t_flush>0,h_t_flush<15),{start}+{slope}*h_t_flush,IF(AND(h_t_flush>=15,h_t_flush<180),IF(AND(h_t_flush>=140,AND(h_t_shed>=1,h_t_shed<=40)),{late}-0.5*(h_t_flush-140+1),IF(h_t_shed>=40,{end},{base}+{mid}*(h_t_flush-15+1))),{end})))"
            a("e_lma_"+tag,expr,"TKY LMA "+tag,"gDM m-2","leaf_aging.c","f_sla_change")
        a("e_t_sla","10000/(c_comp1*e_lma_one+c_comp2*e_lma_two)","当日樹木SLA（Config基準値を上書き）","cm2 gDM-1","leaf_aging.c","f_sla_change")
        for p in ("t","g","v"):
            if p!="t": a(f"e_{p}_sla",f"p_{p}_sla",p+" SLA","cm2 gDM-1")
            a(f"e_{p}_lai",f"MAX(0,e_{p}_sla*x_{p}_fol*c_dmTc/100/2)",p+" 初期LAI","m2 m-2","ecophysiology.c","lai_mass")
            a(f"e_{p}_ek",f"p_{p}_eK0/MAX(0.3,MIN(1,SIN(ASIN(e_sinnoon)*c_rTd*c_dTr)))",p+" 消光係数（原典の角度往復換算）","1","ecophysiology.c","irr_attn")
        a("e_cov_t","1-EXP(-p_t_eK0*e_t_lai)","樹木被覆率")
        a("e_cov_g","(1-e_cov_t)*(1-EXP(-p_g_eK0*e_g_lai))","林床被覆率")
        a("e_cov_v","0","C4被覆率=0 TKY")
        a("e_cov_s","1-e_cov_t-e_cov_g","裸地被覆率")
        a("e_abs_t","1-EXP(-e_t_ek*e_t_lai*0.9)","樹木短波吸収割合")
        a("e_abs_g","(1-e_abs_t)*(1-EXP(-e_g_ek*e_g_lai)*0.9)","林床短波吸収（括弧は原典どおり）")
        a("e_abs_v","0","C4短波吸収=0 TKY")
        a("e_abs_s","1-e_abs_t-e_abs_g","地面短波吸収割合")
        a("e_albsoil","MIN(0.7,MAX(0.05,p_s_albedo0+(0.7-p_s_albedo0)/(1+EXP(-0.05*(x_w_snow-70)))))","積雪を含む地面アルベド","1","soil_physics.c","albedo_soil")
        a("e_lw","((1-z_cloud)*(0.53+0.06*SQRT(z_vp))+z_cloud)*(z_ta+c_ZAT)**4*0.000000056703-9*z_cloud-0.95*(z_ts+c_ZAT)**4*0.000000056703","長波収支（原典符号）","W m-2","radiation.c","f_net_rad")
        for p in ("t","g","v","s"):
            alb="e_albsoil" if p=="s" else f"p_{p}_albedo"
            a(f"e_rn_{p}",f"e_swmean*e_abs_{p}*(1-{alb})-e_lw*e_cov_{p}",p+" 純放射（原典符号）","W m-2","radiation.c","f_net_rad")
        self.issue("SRC07","純放射の長波符号","f_net_radはlw_down-lw_upを定義した後、短波からこれを引く。原典の演算を保持。","意図または修正版")
        # Actual water update, exactly in native order with zero-area C4 retained.
        def w(name,expr,label):
            return self.add(name,expr,label,"mm day-1 または段階貯留mm","21_水過程","hydro_balance.c","f_hydrology")
        for p in ("t","g","v"):
            self.add(f"w_pmi_{p}",f"MAX(0,e_dl*(e_slope*e_rn_{p}+c_cpwet*e_density*z_vpd/e_ra)/(e_slope+0.667)/c_LHT)",p+" PM遮断蒸発","mm day-1","21_水過程","hydro_flows.c","pm_incep")
            self.add(f"w_pmt_{p}",f"IF(AND(h_{p}_gc*c_weight_{p}>0,e_rn_{p}>0),MAX(0,e_dl*(e_slope*e_rn_{p}+c_cpwet*e_density*z_vpd/e_ra)/(e_slope+0.667*(1+1/(h_{p}_gc*c_weight_{p}*c_eta*e_ra)))/c_LHT),0)",p+" PM蒸散（前日gc）","mm day-1","21_水過程","hydro_flows.c","pm_transp")
        self.add("w_rs","1/((1000*x_w_sw/p_fc30+100)*c_eta)","土壌抵抗","s m-1","21_水過程","hydro_flows.c","pm_evap")
        self.add("w_pme","MAX(0,e_dl*(e_slope*e_rn_s+c_cpsoil*e_density*z_vpd/e_ra)/(e_slope+0.667*(1+w_rs/e_ra))/c_LHT)","PM土壌蒸発（cp原典）","mm day-1","21_水過程","hydro_flows.c","pm_evap")
        w("w_fsnow","1/(1+EXP(0.75*(z_ta-2)))","降雪割合（無次元）")
        w("w_snowfall","w_fsnow*z_rain","降雪入力")
        w("w_rainfall","(1-w_fsnow)*z_rain","降雨入力")
        w("w_melt","IF(x_w_snow>0.1,(1/11)/(1+EXP(-0.5*(z_ta-4)))*(1+10/(0.05*x_w_snow+1))*x_w_snow,x_w_snow)","融雪")
        for p in ("t","g","v"):
            rain="w_rainfall" if p=="t" else f"c_weight_{p}*(w_rainfall-w_int_t)"
            w(f"w_capt_{p}",f"MIN({rain},c_weight_{p}*e_{p}_lai*0.25)",p+" 葉面で捕捉可能な雨")
            w(f"w_int_{p}",f"(w_capt_{p}+w_pmi_{p}-SQRT((w_capt_{p}+w_pmi_{p})**2-3.4*w_capt_{p}*w_pmi_{p}))/1.7",p+" 実遮断蒸発")
        w("w_interception","w_int_t+w_int_g+w_int_v","遮断蒸発合計")
        w("w_through","w_rainfall-w_interception","林内雨")
        w("w_liquid","w_through+w_melt","上層への液体流入")
        w("w_ro1","MAX(0,MAX(0,w_liquid**3+(p_fc30-x_w_sw)**3)**0.33333-(p_fc30-x_w_sw))","上層から深層への流出")
        w("w_sw1","x_w_sw+w_liquid-w_ro1","流入後の上層水")
        w("w_evap","MAX(0,(w_sw1+w_pme-SQRT((w_sw1+w_pme)**2-3.4*w_sw1*w_pme))/1.7)","実土壌蒸発")
        w("w_sw2","w_sw1-w_evap","蒸発後上層水")
        for p,stock in [("g","w_sw2"),("v","w_sw2"),("t","x_w_dw")]:
            w(f"w_tr_{p}",f"({stock}+w_pmt_{p}-SQRT(({stock}+w_pmt_{p})**2-3.4*{stock}*w_pmt_{p}))/1.7",p+" 実蒸散")
        w("w_sw3","w_sw2-w_tr_g-w_tr_v","林床蒸散後上層水")
        w("w_base","IF(z_th>0,0.001*x_w_dw,0)","基底流：原典では2回控除")
        w("w_dw1","x_w_dw-w_base","基底流1回目控除後深層水")
        w("w_runoff_bucket","MAX(0,MAX(0,w_ro1**3+(p_fc-w_dw1)**3)**0.33333-(p_fc-w_dw1))","深層バケツ流出")
        w("w_ro2","w_runoff_bucket+w_base","原典出力ro2")
        w("w_retran0","(w_dw1*p_fc30/p_fc-w_sw3)/(1+p_fc30/p_fc)","再配分前の水分勾配")
        w("w_retran","IF(w_retran0>0,MIN(w_retran0*0.5,p_hydcond*86400000),MAX(w_retran0,-p_hydcond*86400000))","層間再配分：正=深層から上層")
        w("w_swclip","MAX(0,-(w_sw3+w_retran))","上層ゼロ丸め補正")
        w("w_dwclip","MAX(0,-(w_dw1-w_retran))","深層ゼロ丸め補正")
        w("w_sw_end","w_sw3+w_retran+w_swclip","最終上層水")
        w("w_dw_end","w_dw1-w_retran+w_dwclip+w_ro1-w_ro2-w_tr_t","最終深層水：終端では丸めなし")
        w("w_snow_end","x_w_snow+w_snowfall-w_melt","最終雪貯留")
        w("w_aet","w_interception+w_tr_t+w_tr_g+w_tr_v+w_evap","実蒸発散")
        w("w_pet","w_pme+w_pmi_t+w_pmi_g+w_pmi_v+w_pmt_t+w_pmt_g+w_pmt_v","原典PET")
        w("w_upward","MAX(0,w_retran)","深層→上層")
        w("w_downward","MAX(0,-w_retran)","上層→深層再配分")
        for nm,don,rec in [("w_snowfall",None,{"x_w_snow":"1"}),("w_through",None,{"x_w_sw":"1"}),
                           ("w_melt","x_w_snow",{"x_w_sw":"1"}),("w_ro1","x_w_sw",{"x_w_dw":"1"}),
                           ("w_evap","x_w_sw",{}),("w_tr_g","x_w_sw",{}),("w_tr_v","x_w_sw",{}),
                           ("w_tr_t","x_w_dw",{}),("w_base","x_w_dw",{}),("w_ro2","x_w_dw",{}),
                           ("w_upward","x_w_dw",{"x_w_sw":"1"}),("w_downward","x_w_sw",{"x_w_dw":"1"}),
                           ("w_swclip",None,{"x_w_sw":"1"}),("w_dwclip",None,{"x_w_dw":"1"})]:
            self.flux(nm,don,rec)
        a("e_appl","MAX(0,MIN(1,(p_fc30-w_sw_end)/p_fc30))","上層空隙率","1","location_proc.c")
        a("e_apph","MAX(0,MIN(1,(p_fc-w_dw_end)/p_fc))","深層空隙率","1","location_proc.c")
        a("e_wfps","MAX(0.05,w_sw_end/300/(1-p_density/2.65))","WFPS（原典では上限なし）","1","location_proc.c")
        a("e_psil","-0.05-0.478*(MAX(0.2,w_sw_end)/p_fc30)**(-5.39)","上層水ポテンシャル（土性1）","原典長さ換算;要単位確認","location_proc.c")
        a("e_psih","-1-0.478*(MAX(0.2,w_dw_end)/p_fc)**(-5.39)","深層水ポテンシャル（土性1）","原典長さ換算;要単位確認","location_proc.c")

    def _physiology_phenology(self):
        for p in ("t","g","v"):
            def a(s,expr,label,unit="1",file="ecophysiology.c",fn="f_ecophysiology"):
                return self.add(f"pht_{p}_{s}",expr,p+":"+label,unit,"22_光合成と生理",file,fn)
            a("ppfd", "e_ppfdc" if p=="t" else "e_ppfdc*(1-e_abs_t)","樹冠頂PPFD","umolPhoton m-2 s-1","radiation.c","f_net_rad")
            a("ci0","0.7*z_co2","葉内CO2初期値","ppmv")
            for j in range(1,7):
                ci=f"pht_{p}_ci{j-1}"
                a(f"lue{j}",f"p_{p}_lue0"+(f"*(52-z_ts)/(3.5+0.75*(52-z_ts))*{ci}/(90+0.6*{ci})" if p!="v" else ""),f"反復{j} 量子収率","mol mol-1","ecophysiology.c","f_photo_qy")
                a(f"opt{j}",f"p_{p}_topt0"+(f"+0.01*{ci}" if p!="v" else ""),f"反復{j} 最適温度","degC","photosynthesis.c","f_pc_sat")
                a(f"comp{j}",f"p_{p}_cmpcd0"+("*MAX(0,1+0.0451*(z_ts-20)+0.000347*(z_ts-20)**2)" if p!="v" else ""),f"反復{j} CO2補償点","ppmv","photosynthesis.c","f_pc_sat")
                a(f"ft{j}",f"IF(z_tl<2,0,MAX(0,MIN(1,(z_ts-p_{p}_tmax)*(z_ts-p_{p}_tmin)/((z_ts-p_{p}_tmax)*(z_ts-p_{p}_tmin)-(z_ts-pht_{p}_opt{j})**2))))",f"反復{j} 温度制限","1","photosynthesis.c","f_pc_sat")
                base,amp=(.3,.7) if p!="v" else (.5,.5)
                a(f"fc{j}",f"MAX(0,MIN(1,{base}+{amp}*({ci}-pht_{p}_comp{j})/(p_{p}_kmci+{ci})))",f"反復{j} CO2制限","1","photosynthesis.c","f_pc_sat")
                base,amp=(.05,.95) if p!="v" else (.11,.89)
                a(f"fw{j}",f"MAX(0,MIN(1,{base}+{amp}*w_dw_end/(w_dw_end+p_fc*p_{p}_km_nstl)))",f"反復{j} 水分制限","1","photosynthesis.c","f_pc_sat")
                a(f"psat{j}",f"p_{p}_pmax*pht_{p}_ft{j}*pht_{p}_fc{j}*pht_{p}_fw{j}",f"反復{j} 飽和光合成","umolCO2 m-2 s-1","photosynthesis.c","f_pc_sat")
                a(f"ptop{j}",f"IF(pht_{p}_psat{j}+pht_{p}_lue{j}*e_ppfd>0,pht_{p}_psat{j}*pht_{p}_lue{j}*e_ppfd/(pht_{p}_psat{j}+pht_{p}_lue{j}*e_ppfd),0)",f"反復{j} 単葉光合成","umolCO2 m-2 s-1")
                a(f"gs{j}",f"IF(pht_{p}_ptop{j}>0,p_{p}_gs_b0+p_{p}_gs_b1*pht_{p}_ptop{j}/((z_co2-pht_{p}_comp{j})*(1+z_vpd/p_{p}_gs_b2)),p_{p}_gs_b0)",f"反復{j} 気孔コンダクタンス","mmolH2O m-2 s-1","ecophysiology.c","f_stom_cond")
                a(f"ci{j}",f"MAX(0,MIN(z_co2,z_co2-pht_{p}_ptop{j}/(pht_{p}_gs{j}/1.56/1000)))",f"反復{j} 葉内CO2","ppmv","ecophysiology.c","f_incelco2")
            a("bb",f"IF(pht_{p}_psat6>0,e_{p}_ek*pht_{p}_lue6*pht_{p}_ppfd/pht_{p}_psat6,0)","樹冠光応答係数")
            a("gc",f"p_{p}_gs_b0*e_{p}_lai+p_{p}_gs_b1/((z_co2-pht_{p}_comp6)*(1+z_vpd/p_{p}_gs_b2))*IF(pht_{p}_psat6>0,2*pht_{p}_psat6/e_{p}_ek*LN((1+SQRT(1+pht_{p}_bb))/(1+SQRT(1+pht_{p}_bb*EXP(-e_{p}_ek*e_{p}_lai)))),0)","樹冠気孔コンダクタンス（翌日への記憶）","mmolH2O m-2 s-1","ecophysiology.c","f_canopy_cond")
            for organ,qkey in [("fol","qTf0"),("stm","qTc0"),("rot","qTr0")]:
                a("q_"+organ,f"p_{p}_{qkey}*EXP(-0.009*(z_ts-15))",organ+" 呼吸Q10","1","ecophysiology.c","f_q10_ar")
            a("sap",f"MIN(x_{p}_stm,x_{p}_stm**(1-0.33334*x_{p}_stm/(p_{p}_f_sz_s+x_{p}_stm)))","辺材炭素","MgC ha-1","ecophysiology.c","f_spcfc_resp")
            a("fine",f"MIN(x_{p}_rot,x_{p}_rot**(1-0.33334*x_{p}_rot/(p_{p}_f_sz_r+x_{p}_rot)))","細根炭素","MgC ha-1","ecophysiology.c","f_spcfc_resp")
            a("rm_stm",f"(p_{p}_rmc_s*pht_{p}_sap+p_{p}_rmc_h*(x_{p}_stm-pht_{p}_sap))/(x_{p}_stm+0.00001)","幹比呼吸係数","10^-3 day-1","ecophysiology.c","f_spcfc_resp")
            a("rm_rot",f"(p_{p}_rmr_s*pht_{p}_fine+p_{p}_rmr_h*(x_{p}_rot-pht_{p}_fine))/(x_{p}_rot+0.00001)","根比呼吸係数","10^-3 day-1","ecophysiology.c","f_spcfc_resp")
            a("nscmax",f"0.1*pht_{p}_sap+0.3*pht_{p}_fine","NSC最大容量","MgC ha-1")
            a("cost",f"(p_{p}_rmf0*EXP(LN(pht_{p}_q_stm)/10*(z_ts-15))/1000+p_{p}_lf0*(1+p_{p}_rgf))*c_dmTc*10000/e_{p}_sla","最適LAI計算用費用","原典換算単位","ecophysiology.c","f_opt_lai")
            a("optlai",f"IF(pht_{p}_psat6>0,LN(MAX(1,e_{p}_ek*pht_{p}_lue6*pht_{p}_ppfd/(pht_{p}_psat6*(pht_{p}_psat6*e_dl/(pht_{p}_psat6*e_dl-pht_{p}_cost*24)-1))))/e_{p}_ek,0)","最適LAI","m2 m-2","ecophysiology.c","f_opt_lai")
            a("narea",f"IF(e_{p}_lai>0.01,x_{p}_ncan/10000/14*1000/e_{p}_lai,1)","葉面積当たり窒素","mmolN m-2leaf","ecophysiology.c","f_n_conc")
            a("nstruct",f"MIN(x_{p}_stm*1000000/250+x_{p}_rot*1000000/150,x_{p}_nstr*0.95)","構造窒素診断","gN ha-1","ecophysiology.c","f_n_conc")
            a("nmobile",f"MAX(1,0.025*(pht_{p}_sap+pht_{p}_fine)*1000000)","可動窒素容量","gN ha-1","ecophysiology.c","f_n_conc")
            # Root depth diagnostic is a finite 60-step traversal, not a circular formula.
            for j in range(1,61):
                a(f"rootlayer{j}",f"x_{p}_rot*100*(0.5*(EXP(-p_{p}_root_dpt_a*{(j-1)/100})+EXP(-p_{p}_root_dpt_b*{(j-1)/100}))-0.5*(EXP(-p_{p}_root_dpt_a*{j/100})+EXP(-p_{p}_root_dpt_b*{j/100})))",f"根深さ探索層{j}","gC m-2")
            expr="0.6"
            for j in range(59,0,-1):
                expr=f"IF(pht_{p}_rootlayer{j}<=0.5,{j/100},{expr})"
            a("rootdepth",expr,"根深さ（60層の有限探索）","m")
            def q(s,expr,label,unit="1"):
                return self.add(f"q_{p}_{s}",expr,p+":"+label,unit,"23_フェノロジー","phenology.c","f_growth_period")
            q("gdd1",f"IF(z_ta<-7,0,h_{p}_gdd+MAX(0,z_ta-p_{p}_ctmp_lfdsp))","日積算後GDD","degC day")
            q("cdd1",f"IF(z_ta>18,0,h_{p}_cdd+IF(AND(z_ta<=p_{p}_ctmp_lfshd,z_doy>=210),z_ta-p_{p}_ctmp_lfshd,0))","日積算後CDD","degC day")
            if p=="t":
                q("season",f"IF(q_{p}_gdd1<260,0,IF(q_{p}_gdd1<580,2,IF(q_{p}_cdd1<(-190),0,IF(q_{p}_cdd1<(-20),3,1))))","季節0休眠/1成長/2展葉/3落葉")
                q("gdd",f"IF(AND(q_{p}_gdd1>=580,q_{p}_cdd1<(-190)),0,q_{p}_gdd1)","翌日GDD","degC day")
                q("cdd",f"IF(q_{p}_season==2,0,q_{p}_cdd1)","翌日CDD","degC day")
                q("flush",f"IF(OR(q_{p}_season==0,q_{p}_season==3),0,h_{p}_flush+1)","翌日展葉後日数","day")
                q("shed",f"IF(q_{p}_season==2,0,IF(OR(q_{p}_season==0,q_{p}_season==3),h_{p}_shed+1,h_{p}_shed))","翌日落葉後日数","day")
            else:
                q("season",f"IF(z_ts>=p_{p}_tmin,1,0)","林床常緑型の成長期")
                q("gdd",f"q_{p}_gdd1","翌日GDD","degC day")
                q("cdd",f"q_{p}_cdd1","翌日CDD","degC day")
                q("flush","0","翌日展葉後日数","day")
                q("shed",f"h_{p}_shed","翌日落葉後日数","day")

    def _plants(self):
        for p in ("t","g","v"):
            def a(s,expr,label,fn="plant_process",file="plant_proc.c",unit="MgC ha-1 day-1 または段階貯留MgC ha-1"):
                return self.add(f"c_{p}_{s}",expr,p+":"+label,unit,"24_植物炭素",file,fn)
            a("emerge",f"IF(q_{p}_season==2,MAX(0,0.03*x_{p}_nsc),0)","NSC展葉への動員")
            a("leaf1",f"x_{p}_fol+0.9*c_{p}_emerge","展葉後の葉量")
            a("lf_rate",f"IF(q_{p}_season==3,IF(e_{p}_lai>0.05,p_{p}_dcd,1),p_{p}_lf0)","葉更新率（Nにも使用）",unit="day-1")
            a("lf",f"c_{p}_leaf1*(c_{p}_lf_rate+IF(AND(q_{p}_season!=3,AND(e_psih<(-148.075),AND(e_{p}_lai>0.05,z_ta>(-5)))),0.012,0))","落葉＋乾燥誘発落葉")
            a("lc",f"p_{p}_lc0*x_{p}_stm","幹更新","f_lc","litterfall.c")
            a("lr",f"p_{p}_lr0*x_{p}_rot","根更新","f_lr","litterfall.c")
            a("leaf2",f"c_{p}_leaf1-c_{p}_lf","落葉後の葉量")
            a("stem2",f"x_{p}_stm-c_{p}_lc","更新後の幹量")
            a("root2",f"x_{p}_rot-c_{p}_lr","更新後の根量")
            a("lai2",f"MAX(0,e_{p}_sla*c_{p}_leaf2*c_dmTc/100/2)","GPP評価時のLAI",unit="m2 m-2")
            a("gpp",f"IF(pht_{p}_psat6>0,2*pht_{p}_psat6*e_dl*(3600*12/100000000)/e_{p}_ek*LN((1+SQRT(1+pht_{p}_bb))/(1+SQRT(1+pht_{p}_bb*EXP(-e_{p}_ek*c_{p}_lai2)))),0)","GPP","f_gpp","photosynthesis.c")
            for organ,stock,rm,temp,fn in [("f","leaf2",f"p_{p}_rmf0","z_ts","f_rfm"),("c","stem2",f"pht_{p}_rm_stm","z_ts","f_rcm"),("r","root2",f"pht_{p}_rm_rot","z_tl","f_rrm")]:
                q={"f":"fol","c":"stm","r":"rot"}[organ]
                a("rm"+organ,f"MAX(0,c_{p}_{stock})*{rm}/1000*EXP(LN(pht_{p}_q_{q})/10*({temp}-15))",organ+" 維持呼吸",fn,"respiration.c")
            a("rpm",f"c_{p}_rmf+c_{p}_rmc+c_{p}_rmr","維持呼吸合計")
            a("epp",f"c_{p}_gpp-c_{p}_rpm","EPP=GPP−維持呼吸")
            a("demand",f"(pht_{p}_optlai-c_{p}_lai2)*100*2/2.2/e_{p}_sla","不足葉量","f_allocation","allocation.c")
            a("af",f"IF(c_{p}_epp>0,IF(c_{p}_lai2>pht_{p}_optlai,0,IF(c_{p}_demand<=c_{p}_epp*p_{p}_alloc_ass,p_{p}_alloc_ass,MIN(c_{p}_demand/c_{p}_epp,0.05))*IF(q_{p}_season==0,0.7,1)),p_{p}_alloc_ass)","葉配分係数","f_allocation","allocation.c","1")
            a("ac",f"(1-c_{p}_af)*p_{p}_alloc_abg","幹配分係数","f_allocation","allocation.c","1")
            a("ar",f"(1-c_{p}_af)*(1-p_{p}_alloc_abg)","根配分係数","f_allocation","allocation.c","1")
            for organ,rg in [("f","rgf"),("c","rgc"),("r","rgr")]:
                a("tp"+organ,f"IF(c_{p}_epp>0,c_{p}_a{organ}*c_{p}_epp,c_{p}_a{organ}*c_{p}_gpp-c_{p}_rm{organ})",organ+" 器官配分（符号付き）","f_allocation","allocation.c")
                a("in"+organ,f"MAX(0,c_{p}_tp{organ})",organ+" 正の純入力 μ")
                a("def"+organ,f"MAX(0,-c_{p}_tp{organ})",organ+" 生産赤字による減少")
                a("rg"+organ,f"IF(c_{p}_epp>0,p_{p}_{rg}*c_{p}_tp{organ}"+(f"+0.1*c_{p}_emerge" if organ=="f" else "")+",0)",organ+" 成長呼吸（展葉費用を含む）")
            a("gatec",f"IF(AND(c_{p}_tpc>0,x_{p}_nsc-c_{p}_emerge<pht_{p}_nscmax),1,0)","幹配分をNSCへ回す条件",unit="0/1")
            a("nsc1",f"x_{p}_nsc-c_{p}_emerge+c_{p}_gatec*c_{p}_inc","幹配分後NSC")
            a("gater",f"IF(AND(c_{p}_tpr>0,c_{p}_nsc1<pht_{p}_nscmax),1,0)","根配分をNSCへ回す条件（幹の後）",unit="0/1")
            a("leaf3",f"c_{p}_leaf2+c_{p}_tpf-c_{p}_rgf","配分後葉量")
            a("stem3",f"c_{p}_stem2+(1-c_{p}_gatec)*c_{p}_inc-c_{p}_defc-c_{p}_rgc","配分後幹量")
            a("root3",f"c_{p}_root2+(1-c_{p}_gater)*c_{p}_inr-c_{p}_defr-c_{p}_rgr","配分後根量")
            a("critfol",f"0.1*100*2/2.2/e_{p}_sla","生存最低葉量（DBF枝）","reallocation_survival","allocation.c")
            for org,stock,factor,alloc in [("sf","stem3",.05,f"p_{p}_alloc_abg"),("rf","root3",.1,f"(1-p_{p}_alloc_abg)")]:
                a("surv"+org,f"IF(c_{p}_leaf3<c_{p}_critfol,c_{p}_critfol*{alloc}*({factor}*c_{p}_{stock}/c_{p}_critfol)/(0.5+{factor}*c_{p}_{stock}/c_{p}_critfol),0)",org+" 生存のため葉へ再配分","reallocation_survival","allocation.c")
            a("stem4",f"c_{p}_stem3-c_{p}_survsf","生存再配分後幹量")
            a("root4",f"c_{p}_root3-c_{p}_survrf","生存再配分後根量")
            a("repair_s",f"IF(AND(c_{p}_stem4<=0,c_{p}_root4>0),c_{p}_root4*p_{p}_alloc_abg,0)","根から幹を再建","reallocation_survival","allocation.c")
            a("repair_r",f"IF(AND(c_{p}_root4-c_{p}_repair_s<=0,c_{p}_stem4+c_{p}_repair_s>0),(c_{p}_stem4+c_{p}_repair_s)*(1-p_{p}_alloc_abg),0)","幹から根を再建","reallocation_survival","allocation.c")
            for organ,expr in [("fol",f"c_{p}_leaf3+c_{p}_survsf+c_{p}_survrf"),("stm",f"c_{p}_stem4+c_{p}_repair_s-c_{p}_repair_r"),("rot",f"c_{p}_root4-c_{p}_repair_s+c_{p}_repair_r"),("nsc",f"c_{p}_nsc1+c_{p}_gater*c_{p}_inr")]:
                a("end_"+organ,expr,"日次写像の最終 "+organ)
            a("npp",f"c_{p}_gpp-c_{p}_rpm-c_{p}_rgf-c_{p}_rgc-c_{p}_rgr","原典NPP")
            self.flux(f"c_{p}_emerge",f"x_{p}_nsc",{f"x_{p}_fol":"0.9"},"NSC動員の10%費用を受取係数に反映")
            for nm,org in [("lf","fol"),("lc","stm"),("lr","rot")]:
                s="t"+nm[-1] if p=="t" else "g"+nm[-1]
                self.flux(f"c_{p}_{nm}",f"x_{p}_{org}",{"x_s_"+s:f"c_weight_{p}"})
            for organ,org,gate in [("f","fol",None),("c","stm",f"c_{p}_gatec"),("r","rot",f"c_{p}_gater")]:
                rec={f"x_{p}_{org}":"1"} if gate is None else {f"x_{p}_{org}":f"1-{gate}",f"x_{p}_nsc":gate}
                self.flux(f"c_{p}_in{organ}",None,rec,"正の器官純入力。維持呼吸を二度引かない。")
                self.flux(f"c_{p}_def{organ}",f"x_{p}_{org}")
                self.flux(f"c_{p}_rg{organ}",f"x_{p}_{org}")
            for nm,don,rec in [("survsf","stm","fol"),("survrf","rot","fol"),("repair_s","rot","stm"),("repair_r","stm","rot")]:
                self.flux(f"c_{p}_{nm}",f"x_{p}_{don}",{f"x_{p}_{rec}":"1"})

    def _soil(self):
        def a(n,e,l,u="1",file="decomposition.c",fn=None):
            return self.add(n,e,l,u,"25_土壌炭素",file,fn)
        for layer,tmp,water,cap,km,ks,app in [("l","z_tl","w_sw_end","p_fc30","kml","kmsl","e_appl"),("h","z_th","w_dw_end","p_fc","kmh","kmsh","e_apph")]:
            a("s_ft"+layer,f"IF({tmp}>(-20),0.01+EXP(308.56*(1/56.02-1/({tmp}+46.02))),0.01)",layer+" 分解温度倍率",fn="frl" if layer=="l" else "frh")
            a("s_fw"+layer,f"MIN(0.8*{water}/(p_s_{km}*{cap}+{water})+0.2,0.4*{app}*p_s_{ks}/(p_s_{ks}+{app})+0.6)",layer+" 水分・通気制限",fn="frl" if layer=="l" else "frh")
            a("s_xi"+layer,f"s_ft{layer}*s_fw{layer}",layer+" 分解倍率 ξ")
        for pool,base in [("tf","lf"),("tc","lc"),("tr","lr"),("gf","lf"),("gc","lc"),("gr","lr"),("ha","ha"),("hi","hi"),("hp","hp")]:
            hum=pool.startswith("h")
            a("s_decomp_"+pool,f"x_s_{pool}*p_s_sr_{base}/1000*s_xi{'h' if hum else 'l'}",pool+" 分解","MgC ha-1 day-1","soil_proc.c","f_cycle_soil")
            a("s_resp_"+pool,"s_decomp_"+pool+("" if hum else f"*p_s_f_co2_{base}"),pool+" 微生物呼吸","MgC ha-1 day-1","soil_proc.c","f_cycle_soil")
            rec={} if hum else {"x_s_"+h:f"(1-p_s_f_co2_{base})*p_s_f_hm_{h[-1]}" for h in ("ha","hi","hp")}
            self.flux("s_decomp_"+pool,"x_s_"+pool,rec)
        for pool in ("tf","tc","tr","gf","gc","gr","ha","hi","hp"):
            terms=[f"-s_decomp_{pool}"]
            for f in self.fluxes:
                if "x_s_"+pool in f.recipients:
                    terms.append(f"{f.name}*({f.recipients['x_s_'+pool]})")
            a("s_raw_"+pool,"x_s_"+pool+"+"+"+".join(terms),pool+" 更新後（丸め前）","MgC ha-1","soil_proc.c","f_cycle_soil")
            a("s_clip_"+pool,"MAX(0,-s_raw_"+pool+")",pool+" ゼロ丸め量","MgC ha-1","soil_proc.c","f_cycle_soil")
            a("s_end_"+pool,"MAX(0,s_raw_"+pool+")",pool+" 最終量","MgC ha-1","soil_proc.c","f_cycle_soil")
            self.flux("s_clip_"+pool,None,{"x_s_"+pool:"1"},"物理流入ではなく負値ゼロ丸めの数値補正")
        a("s_rh","+".join("s_resp_"+p for p in ("tf","tc","tr","gf","gc","gr","ha","hi","hp")),"従属栄養呼吸合計","MgC ha-1 day-1","soil_proc.c","f_cycle_soil")
        a("s_litter_end","+".join("s_end_"+p for p in ("tf","tc","tr","gf","gc","gr")),"N無機化評価時リターC","MgC ha-1","n_flux.c","f_n_mineralz")
        a("s_humus_end","s_end_ha+s_end_hi+s_end_hp","N無機化評価時腐植C","MgC ha-1","n_flux.c","f_n_mineralz")
        a("s_litter_resp","+".join("s_resp_"+p for p in ("tf","tc","tr","gf","gc","gr")),"リター呼吸","MgC ha-1 day-1","n_flux.c","f_n_mineralz")

    def _nitrogen(self):
        def a(n,e,l,fn="f_n_uptake",file="n_flux.c",u="gN ha-1 day-1 または段階貯留gN ha-1"):
            return self.add(n,e,l,u,"26_窒素過程",file,fn)
        a("n_biofix","MAX(0,0.234*w_aet*100-0.172)","共生・非共生N固定（原典年率をそのまま使用）","f_biolfix")
        a("n_roots","x_t_rot+x_g_rot","吸収重み分母（TKY）",u="MgC ha-1")
        a("n_ks","0.90*(w_sw_end/p_fc30)**3+0.1","N拡散倍率",u="1")
        for p in ("t","g","v"):
            a(f"n_{p}_share",f"c_weight_{p}*x_{p}_rot/n_roots",p+" 根量割合：分母0は未定義",u="1")
            a(f"n_{p}_bbb",f"(x_{p}_nstr-pht_{p}_nstruct)/pht_{p}_nmobile",p+" 可動N充足度",u="1")
            a(f"n_{p}_stor",f"IF(AND(n_{p}_bbb>0,n_{p}_bbb<=1),1-n_{p}_bbb/(n_{p}_bbb+pht_{p}_nmobile*0.4),0)",p+" 貯蔵Nによる吸収倍率（単位疑義原典保持）",u="1")
            a(f"n_{p}_fix",f"n_{p}_share*n_biofix",p+" N固定","f_biolfix")
            for species in ("no3","nh4"):
                a(f"n_{p}_up_{species}",f"n_{p}_share*n_{p}_stor*x_n_{species}*0.1*n_ks/(1+n_ks*x_n_{species}/10000)*EXP(0.0693*z_tl)",p+" 吸収 "+species)
            a(f"n_{p}_obtain",f"n_{p}_fix+n_{p}_up_no3+n_{p}_up_nh4",p+" 獲得N","f_n_alloc")
            a(f"n_{p}_demand",f"MAX(0,(SQRT(30*70/(30*0.003))-70-pht_{p}_narea)*14/1000*10000)",p+" 葉N需要","f_n_alloc")
            a(f"n_{p}_acan",f"MIN(n_{p}_obtain,n_{p}_demand)",p+" 獲得Nの葉群配分","f_n_alloc")
            a(f"n_{p}_astr",f"n_{p}_obtain-n_{p}_acan",p+" 獲得Nの貯蔵配分","f_n_alloc")
            a(f"n_{p}_frac",f"IF(n_{p}_obtain>0,n_{p}_acan/n_{p}_obtain,0)",p+" 獲得Nの葉群分率","f_n_alloc",u="1")
            a(f"n_{p}_realloc",f"IF(n_{p}_demand>0,MIN(x_{p}_nstr,n_{p}_demand),0)",p+" 貯蔵から葉へN再配分（clearで毎日0）","f_n_realloc")
            a(f"n_{p}_salvage",f"0.3*c_{p}_lf_rate*x_{p}_ncan",p+" 落葉からのN回収","f_n_abandon_salvage")
            a(f"n_{p}_abcan",f"0.7*c_{p}_lf_rate*x_{p}_ncan",p+" 葉群からリターN","f_n_abandon_salvage")
            a(f"n_{p}_abstr",f"MIN(0.02*x_{p}_nstr,IF(c_{p}_end_stm+c_{p}_end_rot>0,(c_{p}_end_stm*p_{p}_lc0+c_{p}_end_rot*p_{p}_lr0)/(c_{p}_end_stm+c_{p}_end_rot)*x_{p}_nstr,0))",p+" 幹根更新に伴う貯蔵N損失","f_n_abandon_salvage")
            for name,donor in [(f"n_{p}_fix",None),(f"n_{p}_up_no3","x_n_no3"),(f"n_{p}_up_nh4","x_n_nh4")]:
                self.flux(name,donor,{f"x_{p}_ncan":f"n_{p}_frac",f"x_{p}_nstr":f"1-n_{p}_frac"})
            for name,donor,rec in [("realloc","nstr","ncan"),("salvage","ncan","nstr")]:
                self.flux(f"n_{p}_{name}",f"x_{p}_{donor}",{f"x_{p}_{rec}":"1"})
            self.flux(f"n_{p}_abcan",f"x_{p}_ncan",{"x_n_lit":"1"})
            self.flux(f"n_{p}_abstr",f"x_{p}_nstr",{"x_n_lit":"1"})
        a("n_minlit","IF(s_litter_end>0,x_n_lit*s_litter_resp/s_litter_end*0.1,0)","リターN無機化","f_n_mineralz")
        a("n_minhum","IF(s_humus_end>0,x_n_hum*(s_resp_ha+s_resp_hi+s_resp_hp)/s_humus_end*0.1,0)","腐植N無機化","f_n_mineralz")
        a("n_immob","0.01*n_minlit+0.1*n_minhum+0.004*x_n_no3+0.04*x_n_nh4","微生物固定化（原典PSOコメントの現用式）","f_n_immoblz")
        a("n_micdeath","0.80*EXP(LN(2)/10*(z_tl-10))*x_n_mic","微生物N更新","f_n_mcrb_abdn")
        a("n_nh3","x_n_nh4*0.02*MAX(0,5.8/30*10**(p_ph-10))*MAX(0,z_tl**2*(45-z_tl)/(20**2*25))*IF(w_sw_end>0,EXP(18*(-10*(p_fc30/w_sw_end)**5)/(8314*(z_tl+c_ZAT))),0)*17/14","NH3揮散（NH3質量のままN収支へ）","f_nh3_volatilization")
        a("n_leach","MIN(IF(w_sw_end>0,w_ro2*0.1*(x_n_no3/10000)/w_sw_end,0),x_n_no3)*10000","硝酸溶脱（上限の単位原典保持）","f_n_leaching")
        for spec in ("no3","nh4"):
            a("n_conc_"+spec,f"x_n_{spec}/(10000*p_density)",spec+" 土壌濃度","f_n2o_emit_ngas","n2o_emit.c","ugN gsoil-1")
        a("n_wfac","((e_wfps-1.27)/(0.60-1.27))**(2.84*(1.27-0.60)/(0.60-0.0012))*((e_wfps-0.0012)/(0.60-0.0012))**2.84","NGAS1 硝化水分倍率","f_n2o_emit_ngas","n2o_emit.c","1")
        a("n_ngas_n2on","n_wfac*(0.56+ATAN(c_pi*0.45*(p_ph-5))/c_pi)*(-0.06+0.13*EXP(0.07*z_tl))*(3.8+30*(1-EXP(-0.0105*n_conc_nh4)))","NGAS1硝化N2O-N（表示換算前）","f_n2o_emit_ngas","n2o_emit.c")
        a("n_ngas_dn","4.82/14**(16/14**(1.39*e_wfps))*MIN(11000+40000*ATAN(c_pi*0.002*(n_conc_no3-180))/c_pi,24000/(1+200/EXP(0.35*s_rh*1000))-100)","NGAS1脱窒総N（原典負値許容）","f_n2o_emit_ngas","n2o_emit.c")
        a("n_ngas_ratio","1.4/13**(17/13**(2.2*e_wfps))*MIN((1-(0.5+ATAN(c_pi*0.01*(n_conc_no3-190))/c_pi))*25,13+30.78*ATAN(c_pi*0.07*(s_rh*1000-13))/c_pi)","NGAS1 N2/N2O比","f_n2o_emit_ngas","n2o_emit.c","1")
        a("n_nitrif","n_ngas_n2on*100","NH4→NO3硝化","f_n2o_emit_ngas","n2o_emit.c")
        a("n_gas_n2on","n_ngas_n2on*0.1","n_budgetの硝化N2O控除（単位不整合あり）","n_budget","n_budget.c")
        a("n_gas_n2od","n_ngas_dn/(1+n_ngas_ratio)*0.1","n_budgetの脱窒N2O控除","n_budget","n_budget.c")
        a("n_gas_n2","n_ngas_dn/(1+1/n_ngas_ratio)*0.1","n_budgetのN2控除","n_budget","n_budget.c")
        for name,don,rec in [("n_minlit","x_n_lit",{"x_n_nh4":"1"}),("n_minhum","x_n_hum",{"x_n_nh4":"1"}),
                              ("n_immob","x_n_nh4",{"x_n_mic":"1"}),("n_micdeath","x_n_mic",{"x_n_hum":"1"}),
                              ("n_nitrif","x_n_nh4",{"x_n_no3":"1"}),("n_nh3","x_n_nh4",{}),
                              ("n_leach","x_n_no3",{}),("n_gas_n2on","x_n_nh4",{}),
                              ("n_gas_n2od","x_n_no3",{}),("n_gas_n2","x_n_no3",{}),
                              ("z_depo_no3",None,{"x_n_no3":"1"}),("z_depo_nh4",None,{"x_n_nh4":"1"})]:
            self.flux(name,don,rec)
        for state in [s for s in self.states if s.group=="N"]:
            terms=[]
            for f in self.fluxes:
                if f.donor==state.name: terms.append("-"+f.name)
                if state.name in f.recipients: terms.append(f"{f.name}*({f.recipients[state.name]})")
            name=state.name[2:]
            a("n_raw_"+name,state.name+"+"+"+".join(terms),state.label+" 丸め前","n_budget","n_budget.c")
            a("n_clip_"+name,"MAX(0,-n_raw_"+name+")",state.label+" ゼロ丸め量","n_budget","n_budget.c")
            a("n_end_"+name,"MAX(0,n_raw_"+name+")",state.label+" 最終値","n_budget","n_budget.c")
            self.flux("n_clip_"+name,None,{state.name:"1"},"非物理的なゼロ丸め補正を独立表示")

    def _diagnostics(self):
        def a(n,e,l,u="1",file="soil_physics.c",fn=None):
            return self.add(n,e,l,u,"27_補助状態とガス",file,fn)
        a("d_saxton_a","EXP(-4.396-0.0715*p_clay*100-0.0004488*(p_sand*100)**2-0.00004285*(p_sand*100)**2*p_clay*100)*100","Saxton a","原典圧力規約",fn="f_soil_saxton")
        a("d_saxton_b","-3.14-0.00222*(p_clay*100)**2-0.00003484*(p_sand*100)**2*p_clay*100","Saxton b",fn="f_soil_saxton")
        a("d_fc","(33/d_saxton_a)**(1/d_saxton_b)*300","CASA用圃場容水量（Site fc30とは別）","mm",fn="f_soil_saxton")
        a("d_pc","(0.332-0.0007251*p_sand*100+0.1276*LOG10(p_clay*100))*300","CASA用間隙容量","mm",fn="f_soil_saxton")
        a("d_rdr","(1+d_saxton_a)/(1+d_saxton_a*(w_sw_end/300)**d_saxton_b)","CASA乾燥倍率",file="hydro_flows.c",fn="f_casa_moisture")
        a("d_casa_raw","x_casa+(z_rain-w_pet)/1000*IF(z_rain>=w_pet,1,d_rdr)","CASA水分更新前","m3 m-3","hydro_flows.c","f_casa_moisture")
        a("d_casa_end","MAX(0,MIN(d_fc/300,d_casa_raw))","CASA水分翌日状態","m3 m-3","hydro_flows.c","f_casa_moisture")
        a("d_casa_net","(d_casa_end-x_casa)/c_dt","CASAの離散増分を日率化","day-1","hydro_flows.c","f_casa_moisture")
        a("d_ecasa","(z_rain-w_pet)/1000-(d_fc/300-x_casa)","CASA余剰水","m3 m-3","hydro_flows.c","f_casa_moisture")
        a("d_iwcasa","MAX(0,MIN(100,IF(d_ecasa>0,(d_ecasa+d_fc/300)/(d_pc/300)*100,d_casa_end/(d_pc/300)*100)))","CASA水分指数","%","hydro_flows.c","f_casa_moisture")
        self.flux("d_casa_net",None,{"x_casa":"1"},"補助状態の離散式。炭素・窒素・水の質量には加算しない。")
        a("d_doc_temp","(z_tl+z_th)/2","DOC評価地温","degC","doc.c","f_doc_boyer")
        a("d_doc_prod","IF(w_dw_end>0,0.11*10**(0.04*d_doc_temp),0)","DOC濃度生成率","mgC L-1 day-1","doc.c","f_doc_boyer")
        a("d_doc_decay","IF(w_dw_end>0,(1-EXP(-0.002*d_doc_temp))*x_doc,x_doc)","DOC分解（低温で負の原典式）","mgC L-1 day-1","doc.c","f_doc_boyer")
        a("d_doc_leach","IF(w_dw_end>0,w_ro2/w_dw_end*x_doc,0)","DOC流出濃度減少","mgC L-1 day-1","doc.c","f_doc_boyer")
        a("d_doc_raw","x_doc+d_doc_prod-d_doc_decay-d_doc_leach","DOC未丸め更新","mgC L-1","doc.c","f_doc_boyer")
        a("d_doc_clip","MAX(0,-d_doc_raw)","DOCゼロ丸め","mgC L-1","doc.c","f_doc_boyer")
        a("d_doc_end","MAX(0,d_doc_raw)","DOC最終濃度","mgC L-1","doc.c","f_doc_boyer")
        a("d_doc_output","d_doc_end*w_ro2*10","DOC流出診断（NECB_DOC=0）","gC ha-1 day-1","doc.c","f_doc_boyer")
        for name,donor in [("d_doc_prod",None),("d_doc_decay","x_doc"),("d_doc_leach","x_doc"),("d_doc_clip",None)]:
            self.flux(name,donor,{} if donor else {"x_doc":"1"})
        a("d_gpp","c_t_gpp+c_g_gpp","生態系GPP（C4重み0）","MgC ha-1 day-1","daily_scheme.c","daily_scheme")
        a("d_npp","c_t_npp+c_g_npp","生態系NPP","MgC ha-1 day-1","daily_scheme.c","daily_scheme")
        a("d_nep","d_npp-s_rh","生態系NEP","MgC ha-1 day-1","daily_scheme.c","daily_scheme")
        a("d_ra","c_t_rpm+c_g_rpm+c_t_rgf+c_g_rgf+c_t_rgc+c_g_rgc+c_t_rgr+c_g_rgr","独立栄養呼吸","MgC ha-1 day-1","daily_scheme.c","daily_scheme")
        a("d_voc_fol","(c_t_end_fol+c_g_end_fol)*100*c_dmTc","VOC用葉乾物面密度","gDM m-2","vocemit.c","f_voc_emit_guenther97")
        a("d_voc_light","0.0027*1.066*e_ppfd/SQRT(1+0.0027**2*e_ppfd**2)*0.5","VOC光倍率",file="vocemit.c",fn="f_voc_emit_guenther97")
        a("d_voc_isot","EXP(95000*(z_ts+c_ZAT-303.15)/(8.314*(z_ts+c_ZAT)*303.15))/(0.961+EXP(230000*(z_ts+c_ZAT-314)/(8.314*(z_ts+c_ZAT)*303.15)))","isoprene温度倍率",file="vocemit.c",fn="f_voc_emit_guenther97")
        a("d_voc_other_t","EXP(0.09*(z_ts+c_ZAT-303.15))","他VOC温度倍率",file="vocemit.c",fn="f_voc_emit_guenther97")
        a("d_voc_age","IF(q_t_flush<=0,0.1,IF(q_t_flush<=30,0.3,IF(q_t_flush<=100,1.1,IF(q_t_flush<=200,0.4,0.1))))","VOC葉齢倍率",file="vocemit.c",fn="f_voc_emit_guenther97")
        for name,pot in [("isopr",8),("monotrp",2.4),("methanl",1.8),("acetone",.87),("actaldhd",.3),("frmardhd",.2),("formacd",.03),("acetacd",.006),("co",.3)]:
            a("d_voc_"+name,f"{pot}*d_voc_fol*e_dl*0.01*d_voc_age*"+("d_voc_light*d_voc_isot" if name=="isopr" else "d_voc_other_t"),name+"放出（NECB_BVOC=0、診断）","gC ha-1 day-1","vocemit.c","f_voc_emit_guenther97")
        a("d_sunshine","MAX(0,e_dl*(1-z_cloud))","直達日照時間","h","ch4_emit.c","f_ch4emit_plant")
        for p in ("t","g","v"):
            a(f"d_{p}_ch4mass",f"IF(q_{p}_season!=0,c_{p}_end_fol*c_dmTc*1000*(d_sunshine*374+(24-d_sunshine)*119)*10**(-10),0)",p+"植物CH4葉量式（診断）","gCH4 m-2 day-1","ch4_emit.c","f_ch4emit_plant")
            a(f"d_{p}_ch4photo",f"IF(AND(d_sunshine>0,c_{p}_npp>0),2*(16/12)*(c_{p}_npp*1000)/30000*(1+(24-d_sunshine)/d_sunshine*119/374),0)",p+"植物CH4生産式（診断）","原典出力単位","ch4_emit.c","f_ch4emit_plant")
        a("d_phi","1-p_density/2.65","間隙率",file="ch4_oxy.c",fn="f_ch4oxy_ridgewell")
        a("d_airp","MAX(0,d_phi-w_sw_end/300)","空気充填間隙率",file="ch4_oxy.c",fn="f_ch4oxy_ridgewell")
        a("d_ridg_k","0.00087*MIN(1,(z_rain+w_sw_end)/w_pet)*IF(z_tl<0,0,EXP(0.0693*z_tl-0.000000856*z_tl**4))","Ridgwell反応係数","原典係数","ch4_oxy.c","f_ch4oxy_ridgewell")
        a("d_ridg_diff","d_phi**(4/3)*(d_airp/d_phi)**(1.5+3/d_saxton_b)*(1+0.0055*z_tl)*0.196","Ridgwell拡散係数","cm2 s-1","ch4_oxy.c","f_ch4oxy_ridgewell")
        a("d_ridg_ch4","z_ch4*d_ridg_diff/30*IF(d_ridg_diff+d_ridg_k*30>0,1-d_ridg_diff/(d_ridg_diff+d_ridg_k*30),1)*616.9","Ridgwell CH4酸化（診断）","mgCH4 m-2 day-1","ch4_oxy.c","f_ch4oxy_ridgewell")
        a("d_curry_b","15.9*p_clay+2.91","Curry保水曲線係数",file="ch4_oxy.c",fn="f_ch4oxy_curry")
        a("d_curry_ps","10**(-1.31*p_sand+1.88)/10*(1/d_phi)**(-d_curry_b)","Curry水ポテンシャル（原典frac_water=1）","kPa","ch4_oxy.c","f_ch4oxy_curry")
        # Expose the undefined source branch with an unassigned input instead of
        # silently repairing 0.2>=ps to 0.2<=ps.
        self.graph.add("z_curry_undefined","Curry未定義区間のr_sm","1","04_気象入力",source=self.src("ch4_oxy.c","f_ch4oxy_curry","0.2 >= ps"),note="通常の気象入力ではない。正式な修正が必要。")
        a("d_curry_sm","IF(d_curry_ps<0.2,1,IF(AND(0.2>=d_curry_ps,d_curry_ps<=100),(1-(LOG10(d_curry_ps)-LOG10(0.2))/(LOG10(100)-LOG10(0.2)))**0.8,IF(d_curry_ps>100,0,z_curry_undefined)))","Curry水分倍率（未定義分岐を保持）",file="ch4_oxy.c",fn="f_ch4oxy_curry")
        a("d_curry_t","IF(OR(z_tl<(-10),z_tl>=43.3),0,IF(z_tl<0,(0.1*z_tl+1)**2,EXP(0.0693*z_tl-0.000000865*z_tl**4)))","Curry温度倍率",file="ch4_oxy.c",fn="f_ch4oxy_curry")
        a("d_curry_ch4","586.7*z_ch4*SQRT(0.196*(1+0.0055*z_tl)*d_phi**(4/3)*(d_airp/d_phi)**(1.5+3/d_curry_b)*0.00005*d_curry_t*d_curry_sm)","Curry CH4酸化（診断）","mgCH4 m-2 day-1","ch4_oxy.c","f_ch4oxy_curry")
        a("d_dg_ch4","(40-18.3*p_density)*MAX(0.1,((10*e_wfps-0.5)/(1.84-0.5))**0.13*((10*e_wfps-55)/(1.84-55))**(0.13*(55-1.84)/(1.84-0.5)))*IF(z_tl<0,0,0.0209*z_tl+0.845)*0.1*16/12","Del Grosso DBF CH4酸化（診断）","mgCH4 m-2 day-1","ch4_oxy.c","f_ch4oxy_delgrosso")
        for tag in ("casa","ngas2"):
            file,fn=("ch4_oxy.c","f_ch4oxy_casa") if tag=="casa" else ("n2o_emit.c","f_n2o_emit_ngas_2")
            def b(s,e,l): return a(f"d_{tag}_{s}",e,tag+":"+l,file=file,fn=fn)
            b("aa","d_fc/300","内間隙容積")
            b("pp","(d_pc-d_fc)/300","間間隙容積")
            b("ss","1-d_pc/300","固相容積")
            b("thp",f"IF(d_casa_end<=d_{tag}_aa,0,MIN(d_{tag}_pp,d_ecasa))","間間隙水分")
            b("aaa",f"d_{tag}_pp-d_{tag}_thp","間間隙気相")
            b("bbb",f"(d_{tag}_aa-d_casa_end)/(d_{tag}_aa+d_{tag}_ss)","正規化内間隙気相")
            for suf,poly,last in [("xx",f"d_{tag}_pp",f"d_{tag}_pp"),("yy",f"d_{tag}_aaa",f"d_{tag}_aaa"),("zz",f"d_{tag}_bbb",f"d_{tag}_aaa")]:
                b(suf,f"0.477*{poly}**3-0.596*{poly}**2+0.437*{last}+0.564",suf+"経験多項式")
            den=f"d_{tag}_aa" if tag=="casa" else "z_ngas2_aa"
            b("jjj",f"d_{tag}_aa-d_casa_end/{den}+d_{tag}_ss","原典jjj（NGAS2 aa未初期化）")
            b("fff",f"(1-d_casa_end/d_{tag}_aa)**2","飽和補正")
            b("ccc",f"d_{tag}_fff*d_{tag}_jjj**(2*d_{tag}_zz)*(1-d_{tag}_pp**(2*d_{tag}_xx))*(d_{tag}_aaa-d_{tag}_aaa**(2*d_{tag}_yy))","拡散式分子")
            b("ddd",f"d_{tag}_fff*d_{tag}_jjj**2*(1-d_{tag}_pp**(2*d_{tag}_xx))+d_{tag}_aaa-d_{tag}_aaa**(2*d_{tag}_yy)","拡散式分母")
            b("eee",f"(1-d_{tag}_thp/d_{tag}_pp)**2*d_{tag}_aaa**(2*d_{tag}_yy)","間間隙拡散")
            b("diff",f"IF(d_{tag}_jjj>0,d_{tag}_ccc/d_{tag}_ddd+d_{tag}_eee,d_{tag}_eee)","正規化拡散係数")
        a("d_casa_ch4","IF(z_tl>0,0.194*(0.9734+0.0055*z_tl)*d_casa_diff*0.04*0.00027*30*24*3600,0)","CASA CH4酸化（診断）","mgCH4 m-2 day-1","ch4_oxy.c","f_ch4oxy_casa")
        a("d_nminnet","MAX(0,n_minlit+n_minhum-n_immob)","CASA/NGAS2正の純無機化","gN ha-1 day-1","n2o_emit.c","f_n2o_emit_casa")
        for gas,coefs in [("no",["0.1*d_iwcasa","-0.005*d_iwcasa+1.05","-0.01*d_iwcasa+1.2","-0.04*d_iwcasa+3","-0.02*d_iwcasa+1.6","0","0"]),("n2o",["0","0.005*d_iwcasa-0.05","0.01*d_iwcasa-0.2","0.04*d_iwcasa-2","0.02*d_iwcasa-0.6","-0.02*d_iwcasa+2.6","-0.08*d_iwcasa+8"]),("n2",["0","0","0","0","0","0.02*d_iwcasa-1.6","0.08*d_iwcasa-7"])]:
            expr=coefs[-1]
            for edge,coef in reversed(list(zip([10,30,60,70,80,90],coefs[:-1]))): expr=f"IF(d_iwcasa<{edge},{coef},{expr})"
            conv={"no":"30/14","n2o":"44/28","n2":"1"}[gas]
            a("d_casa_"+gas,f"({expr})*d_nminnet*0.01*0.1*({conv})","CASA "+gas+"（診断、n_budgetへ入らない）","mg gas m-2 day-1","n2o_emit.c","f_n2o_emit_casa")
        def ng(s,e,l,u="1"): return a("d_ng2_"+s,e,"NGAS2:"+l,u,"n2o_emit.c","f_n2o_emit_ngas_2")
        ng("tmin","IF(z_tsoil_mean>0,9.98*LN(z_tsoil_mean)-13.251,0)","下限地温","degC")
        ng("tmax","IF(z_tsoil_mean>0,6.4642*LN(z_tsoil_mean)+34.215,0)","上限地温","degC")
        ng("topt","IF(z_tsoil_mean>0,6.6804*LN(z_tsoil_mean)+13.926,0)","最適地温","degC")
        ng("ft","IF(z_tsoil_mean>0,MAX(0,(z_tl-d_ng2_tmin)*(z_tl-d_ng2_tmax)/((z_tl-d_ng2_tmin)*(z_tl-d_ng2_tmax)-(z_tl-d_ng2_topt)**2)),0)","温度倍率")
        ng("fn2o","0.02*(d_nminnet/10000)*0.2*0.1*MAX(0,n_conc_nh4*10**(-6)*p_density*3/(0.025**2*c_pi))*d_ng2_ft*MAX(0,5.63*e_wfps-4.64*e_wfps**2-0.745)*MAX(0,0.56+ATAN(c_pi*0.45*(p_ph-5))/c_pi)*10000","硝化N2O-N","gN ha-1 day-1")
        ng("fdw","MAX(0,0.5*ATAN(0.6*c_pi*(10*e_wfps-(0.90-(MIN(0.113,d_ngas2_diff)*(-3.05)+0.36))))/c_pi)","脱窒水分倍率")
        ng("denit","d_ng2_fdw*MIN(1.15*n_conc_no3**0.57,0.1*(s_rh*100/p_density)**1.3)","脱窒総N","原典gN ha-1 day-1")
        ng("kone","MAX(1.7,38.4-350*d_ngas2_diff)","N2/N2O係数")
        ng("ratio","MAX(0.1,1.5*e_wfps-0.32)*MIN(0.16*d_ng2_kone,d_ng2_kone*EXP(-0.8*n_conc_no3/(s_rh*1000)))","N2/N2O比")
        ng("noxratio","15.2+35.5*ATAN(0.68*c_pi*(10*d_ngas2_diff-1.86))/c_pi","NOx比")
        ng("nox","d_ng2_noxratio*(d_ng2_denit/(1+d_ng2_ratio)+d_ng2_fn2o)","NOx原典ローカル診断（状態収支へ入らない）","gN ha-1 day-1")
        ng("n2o","(d_ng2_fn2o+d_ng2_denit/(1+d_ng2_ratio))*44/28*0.1","総N2O（診断）","mgN2O m-2 day-1")
        ng("n2","d_ng2_denit/(1+1/d_ng2_ratio)*0.1","N2（診断）","mgN2 m-2 day-1")
        self.graph.add('z_rain_annual','RUSLE評価年の年降水量','mm year-1','04_気象入力',
                       source=self.src('erosion.c','f_erosion_rusle'),note='年次診断のみ。日次Xの更新には不要。')
        a('d_ero_r','MAX(0,IF(z_rain_annual<=850,0.0483*z_rain_annual**1.610,587.8-1.219*z_rain_annual+0.004105*z_rain_annual**2))','RUSLE降雨係数','原典RUSLE規約','erosion.c','f_erosion_rusle')
        a('d_ero_f','SIN(10*c_dTr)/0.0896/(3*SIN(10*c_dTr)**0.8+0.56)','RUSLE斜面指数の中間量','1','erosion.c','f_erosion_rusle')
        a('d_ero_ls','(16.8*SIN(10*c_dTr)-0.5)*(100/22.13)**(d_ero_f/(1+d_ero_f))','RUSLE長さ・斜面倍率（原典10度、100m）','1','erosion.c','f_erosion_rusle')
        a('d_ero_dg','-3.5*p_sand-2*(1-p_sand-p_clay)-0.5*p_clay','RUSLE土粒径指数','原典経験指数','erosion.c','f_erosion_rusle')
        a('d_ero_k','IF(p_clay>0,0.0293*(0.65-d_ero_dg+0.24*d_ero_dg**2)*EXP(-0.0021/p_clay-0.00037/p_clay**2-4.02*p_clay+1.72*p_clay**2),0)','RUSLE侵食性係数（原典SOC=1%）','原典RUSLE規約','erosion.c','f_erosion_rusle')
        a('d_ero_soil','d_ero_r*d_ero_ls*d_ero_k*0.001','森林RUSLE土壌侵食（年次診断）','Mgsoil ha-1 year-1','erosion.c','f_erosion_rusle')
        a('d_ero_c','d_ero_soil*0.01*0.45','RUSLE炭素搬出（NECB_POC=0、年次診断）','MgC ha-1 year-1','erosion.c','f_erosion_rusle')
        self.issue("SRC08","原典の分数累乗・ゼロ除算","WFPS>1.27、負の器官炭素、PET=0等で原典の実数式が定義されない場合がある。値を丸めて補修しない。","想定する状態・気象の有効範囲、必要なら正式な境界処理")

    def _metadata(self):
        """Resolve bundled source expressions to function definitions, not calls."""
        for name,node in self.graph.nodes.items():
            if node.sheet=='20_気象放射' and '::' not in node.source:
                if name.startswith(('e_cov_','e_app','e_wfps','e_psi')):
                    node.source=self.src('location_proc.c','f_loct_proc')
                elif name.startswith('e_abs_'):
                    node.source=self.src('radiation.c','f_net_rad')
                elif name=='e_ge':node.source=self.src('radiation.c','f_soldec')
                elif name=='e_sinnoon':node.source=self.src('radiation.c','f_solhgt')
                elif name=='e_dtc':node.source=self.src('radiation.c','f_toprad')
                elif name.endswith('_sla'):node.source=self.src('parameter.c','set_parameter')
                else:node.source=self.src('radiation.c','f_ppfd')
        for p in ('t','g','v'):
            self.graph.nodes[f'c_{p}_lf_rate'].note+=' 基準率はecophysiology.c::f_mortality。季節分岐はplant_proc.c。'
            for key,fn in [('rgf','f_rfg'),('rgc','f_rcg'),('rgr','f_rrg')]:
                self.graph.nodes[f'c_{p}_{key}'].source=self.src('respiration.c',fn)
                self.graph.nodes[f'c_{p}_{key}'].note+=' 葉の展葉費用加算はplant_proc.c::plant_process。'
            self.graph.nodes[f'q_{p}_season'].source=self.src('phenology.c','phenology_colddeciduous' if p=='t' else 'phenology_evergreen')
        for name in ('n_gas_n2on','n_gas_n2od','n_gas_n2'):
            self.graph.nodes[name].unit='原典数値: mgN m-2 day-1をgN ha-1 day-1として控除'
        self.graph.nodes['n_biofix'].unit='原典: kgN ha-1 year-1（収支へ無換算）'
        self.graph.nodes['d_doc_clip'].unit='mgC L-1 day-1（1日補正）'
        for name in ('e_psil','e_psih'):
            self.graph.nodes[name].unit='推定: 水頭m（正式単位は要確認）'
            self.graph.nodes[name].note='数値係数と重力項は原典。正式な単位注記は原典コメントで要確認。'
        self.issue('UNIT01','土壌水ポテンシャルの正式単位','係数と重力項は水頭mに対応すると解釈できるが、該当コードに単位の明記がない。','原著・作者の単位規約。圧力単位へ無断換算しない。')
        self.issue('INPUT05','RUSLE年次診断の年降水量','Configに評価年の降水量はない。斜面10度、長さ100m、SOC1%は原典内部定数でありTKY観測値とは限らない。','z_rain_annualを提供。NECB_POC=0のため日次状態へ加えない。')
        self.issue('UNIT02','経験係数と中間診断の単位','gs_b1、最適LAIのcost、RUSLE係数などは原典の混合単位を使う。90の単位列で「原典規約」「混合」と明示。','原著の定義または作者の単位説明。無次元と勝手に扱わない。')
        for node in self.graph.nodes.values():
            if node.name.startswith('z_') and node.name not in ('z_ngas2_aa','z_curry_undefined'):
                node.note+=' 気象はサイト補正等を終えたloct相当の日値。'
            if node.name.endswith(('_phenoltype','_phototype')) or node.name=='p_texture':
                node.note+=' 固定した分類コード。この値だけを変更しても式の分岐は切り替わらない。'

    def _matrix_nodes(self):
        # All internal donor processes are factored by their original-day stock.
        # The zero-donor/cancelling-flux residual restores the exact daily map.
        for state in self.states:
            suffix=state.name[2:]
            fs=[f for f in self.fluxes if f.donor==state.name]
            total="+".join(f.name for f in fs) or "0"
            self.graph.add("m_loss_"+suffix,state.label+" 総控除フラックス",state.unit+" day-1","30_行列組立",expression=total,source=state.source)
            self.graph.add("m_k_"+suffix,state.label+" K対角", "day-1","30_行列組立",expression=state.base,source=state.source,
                           note="C更新・分解は原典基準率。N/W/Dの1 day^-1は因数分解の単位基準で生物パラメータではない。")
            self.graph.add("m_regular_"+suffix,state.label+" 通常因数分解可否","0/1","30_行列組立",
                           expression=f"IF(AND({state.name}!=0,m_loss_{suffix}!=0,m_k_{suffix}!=0),1,0)",source=state.source)
            self.graph.add("m_xi_"+suffix,state.label+" ξ対角","1","30_行列組立",
                           expression=f"IF(m_regular_{suffix}==1,m_loss_{suffix}/(m_k_{suffix}*{state.name}),0)",source=state.source,
                           note="複数過程の合成倍率。環境効果だけでなく日内段階・状態依存を含む。")
        self.external=[f for f in self.fluxes if f.donor is None]
        for state in self.states:
            suffix=state.name[2:]
            raw=[]
            residual=[]
            for f in self.fluxes:
                coeff=[]
                if f.donor==state.name: coeff.append("-1")
                if state.name in f.recipients: coeff.append(f.recipients[state.name])
                if not coeff: continue
                term=f"({'+'.join(coeff)})*{f.name}"
                raw.append(term)
                if f.donor:
                    residual.append(f"IF(m_regular_{f.donor[2:]}==0,{term},0)")
            self.graph.add("m_res_"+suffix,state.label+" ゼロ供与元残差",state.unit+" day-1","30_行列組立",
                           expression="+".join(residual) or "0",source="本ブックの厳密代数分解",
                           note="X_j=0の同日流入後流出、または総控除相殺をBμへ明示的に戻す。")
            self.graph.add("m_rhs_"+suffix,state.label+" 全フラックスの直接和",state.unit+" day-1","30_行列組立",
                           expression="+".join(raw) or "0",source="全フラックス台帳の入出力和")
            self.graph.add("m_end_"+suffix,state.label+" 更新先",state.unit,"30_行列組立",
                           expression=f"{state.name}+c_dt*m_rhs_{suffix}",source="X_next = X + Δt fΔ(X,z,h)")
        self.a_expr={}
        for i,recipient in enumerate(self.states):
            for j,donor in enumerate(self.states):
                if i==j:
                    expr="-1"
                else:
                    terms=[f"{f.name}*({f.recipients[recipient.name]})" for f in self.fluxes if f.donor==donor.name and recipient.name in f.recipients]
                    expr=f"IF(m_regular_{donor.name[2:]}==1,({'+'.join(terms)})/m_loss_{donor.name[2:]},0)" if terms else "0"
                self.a_expr[i,j]=expr
        # Unit annotations for dimensionless/intermediate water cells.
        for n in self.graph.nodes.values():
            if n.name=="w_fsnow": n.unit="1"
            elif n.name.startswith("w_") and (n.name.endswith("_end") or n.name in ("w_sw1","w_sw2","w_sw3","w_dw1","w_retran0")):
                n.unit="mm"
            elif n.sheet=="21_水過程" and n.name!="w_rs": n.unit="mm day-1" if n.name!="w_fsnow" else "1"
            elif n.sheet=="24_植物炭素" and ("end_" in n.name or any(n.name.endswith(s) for s in ("leaf1","leaf2","leaf3","stem2","stem3","stem4","root2","root3","root4","nsc1","critfol","demand"))):
                n.unit="MgC ha-1"
            elif n.unit.startswith("MgC ha-1 day-1 または"): n.unit="MgC ha-1 day-1"
            elif n.unit.startswith("gN ha-1 day-1 または"):
                n.unit="gN ha-1" if n.name.startswith(("n_end_","n_raw_","n_clip_")) else "gN ha-1 day-1（原典数値規約）"
        def budget(name,expr,label,unit,note=''):
            self.graph.add('b_'+name,label,unit,'44_収支検査',expression=expr,
                           source='全フラックスの和。元コードの出力定義との差を明示。',note=note)
        carbon=[s for s in self.states if s.group=='C' and not s.name.startswith('x_v_')]
        nitrogen=[s for s in self.states if s.group=='N']
        budget('c_rate','+'.join('m_rhs_'+s.name[2:] for s in carbon),'実際の生態系C変化率（NSC込み・C4除外）','MgC ha-1 day-1')
        budget('c_clip','+'.join('s_clip_'+p for p in ('tf','tc','tr','gf','gc','gr','ha','hi','hp')),'土壌Cの数値補正流入','MgC ha-1 day-1')
        budget('c_expected','d_nep-0.1*(c_t_emerge+c_g_emerge)+b_c_clip','NEPから展葉の別控除・丸めを補正したC変化','MgC ha-1 day-1')
        budget('c_error','b_c_rate-b_c_expected','C収支残差（0になるべき）','MgC ha-1 day-1')
        budget('c_vs_nep','b_c_rate-d_nep','実C変化−原典NEP（展葉時には0でない）','MgC ha-1 day-1','NEPと全貯留変化を無条件に同一視しない。')
        budget('n_rate','+'.join('m_rhs_'+s.name[2:] for s in nitrogen),'N状態の数値上の合計変化率','gN ha-1 day-1（原典規約）')
        budget('n_clip','+'.join('n_clip_'+s.name[2:] for s in nitrogen),'Nの数値補正流入','gN ha-1 day-1（原典規約）')
        budget('n_expected','n_t_fix+n_g_fix+n_v_fix+z_depo_no3+z_depo_nh4-n_nh3-n_leach-n_gas_n2on-n_gas_n2od-n_gas_n2+b_n_clip','原典N入力−出力＋丸め','gN ha-1 day-1（原典規約）')
        budget('n_error','b_n_rate-b_n_expected','N数値収支残差（単位問題を解決した意味ではない）','gN ha-1 day-1（原典規約）')
        budget('w_rate','m_rhs_w_snow+m_rhs_w_sw+m_rhs_w_dw','雪＋2層水の変化率','mm day-1')
        budget('w_expected','z_rain-w_aet-w_ro2-w_base+w_swclip+w_dwclip','P−AET−ro2−追加baseflow＋丸め','mm day-1')
        budget('w_error','b_w_rate-b_w_expected','原典水収支の数値残差','mm day-1')
        budget('w_vs_reported','b_w_rate-(z_rain-w_aet-w_ro2)','通常のP−AET−ro2との差（追加baseflow等）','mm day-1')

    def matrix_values(self, values):
        """Independent numerical assembly of A, Xi, K, B, mu for verification."""
        n=len(self.states)
        index={s.name:i for i,s in enumerate(self.states)}
        x=np.array([values[s.name] for s in self.states],dtype=float)
        k=np.diag([values["m_k_"+s.name[2:]] for s in self.states])
        xi=np.diag([values["m_xi_"+s.name[2:]] for s in self.states])
        A=-np.eye(n)
        # Expressions for coefficients are intentionally simple here.
        def simple(e):
            t=ast.parse(e,mode="eval")
            allowed={name:v for name,v in values.items() if isinstance(v,(int,float))}
            return eval(compile(t,"<internal coefficient>","eval"),{"__builtins__":{}},allowed)
        for j,s in enumerate(self.states):
            if not values["m_regular_"+s.name[2:]]: continue
            total=values["m_loss_"+s.name[2:]]
            for f in self.fluxes:
                if f.donor==s.name:
                    for r,c in f.recipients.items():
                        A[index[r],j]+=values[f.name]*simple(c)/total
        ext=self.external
        B=np.zeros((n,len(ext)+n))
        mu=[]
        for j,f in enumerate(ext):
            mu.append(values[f.name])
            for r,c in f.recipients.items(): B[index[r],j]=simple(c)
        for i,s in enumerate(self.states):
            B[i,len(ext)+i]=1
            mu.append(values["m_res_"+s.name[2:]])
        return x,B,np.asarray(mu),A,xi,k
