"""Build the auditable TKY matrix workbook, including saved formula results.

Only explicit inputs from Config_TKY and the pinned source receive defaults.
The Python evaluator is not an Excel application; that limitation is recorded.
"""
from __future__ import annotations

import argparse
import ast
import hashlib
import json
import math
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from .visit_workbook import TKYModel, SOURCE_DEFAULT, SOURCE_COMMIT

NAVY = '203B53'
TEAL = '167D8D'
PALE = 'E9F4F5'
AMBER = 'FFF1CE'
BLUE = '1663A6'
GROUP = {'C':'E4F0DF', 'N':'EDE4F5', 'W':'DFEFF8', 'D':'F5E9D7'}
UNKNOWN = '未設定/定義域'


def prepare_matrix(model):
    """Add actual matrix products as DAG nodes, not merely a picture of them."""
    g = model.graph
    n = len(model.states)
    mu = [f.name for f in model.external] + ['m_res_'+s.name[2:] for s in model.states]
    for (i,j),expression in model.a_expr.items():
        g.add(f'mx_a_{i}_{j}',f'A[{i+1},{j+1}]','受取単位/供与単位','13_A',
              expression=expression,source='フラックス台帳を供与状態別に集約',
              note='列=供与元、行=受取先。対角-1。X=0または控除総和0の残差はBμへ。')
    for i,s in enumerate(model.states):
        for j in range(len(mu)):
            expression = (model.external[j].recipients.get(s.name,'0') if j<len(model.external)
                          else str(int(j-len(model.external)==i)))
            g.add(f'mx_b_{i}_{j}',f'B[{i+1},{j+1}]','行の状態単位/μ列の単位','11_B',
                  expression=expression,source='外部・有効入力とゼロ供与元補正の行先')
    for i,s in enumerate(model.states):
        suffix=s.name[2:]
        terms_b=[f'mx_b_{i}_{j}*{mu[j]}' for j in range(len(mu))
                 if g.nodes[f'mx_b_{i}_{j}'].expression!='0']
        terms_m=[f'mx_a_{i}_{j}*m_xi_{d.name[2:]}*m_k_{d.name[2:]}*{d.name}'
                 for j,d in enumerate(model.states) if model.a_expr[i,j]!='0']
        for name,expr,label in [('input','+'.join(terms_b) or '0','Bμ'),
                                ('transfer','+'.join(terms_m) or '0','AξKX'),
                                ('rhs',f'mx_input_{suffix}+mx_transfer_{suffix}','行列右辺'),
                                ('error',f'mx_rhs_{suffix}-m_rhs_{suffix}','行列右辺−直接フラックス和')]:
            g.add(f'mx_{name}_{suffix}',s.label+' '+label,s.unit+' day-1','17_右辺と更新',
                  expression=expr,source='本ブックの行列積（代数恒等式）')
    model.mu_names=mu


def save_with_caches(book, path, caches):
    """Insert OOXML cached results while preserving the real Excel formulas."""
    import io
    buf=io.BytesIO()
    book.save(buf)
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('',ns)
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as zout:
        for entry in zin.infolist():
            data=zin.read(entry.filename)
            match=re.fullmatch(r'xl/worksheets/sheet(\d+)\.xml',entry.filename)
            if match:
                title=book.sheetnames[int(match.group(1))-1]
                tree=ET.fromstring(data)
                for cell in tree.iter('{'+ns+'}c'):
                    if cell.find('{'+ns+'}f') is None: continue
                    key=(title,cell.attrib['r'])
                    if key not in caches: raise ValueError(f'missing formula cache: {key}')
                    value=caches[key]
                    v=cell.find('{'+ns+'}v')
                    if v is None: v=ET.SubElement(cell,'{'+ns+'}v')
                    if isinstance(value,(int,float)):
                        cell.attrib.pop('t',None)
                        v.text=repr(float(value))
                    else:
                        cell.set('t','str')
                        v.text=UNKNOWN if value is None else str(value)
                data=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
            zout.writestr(entry,data)


class Exporter:
    def __init__(self,model):
        self.model=model
        prepare_matrix(model)
        self.values=model.graph.evaluate()
        self.book=Workbook()
        self.book.remove(self.book.active)
        self.book.calculation=CalcProperties(calcId=191029,calcMode='auto',
                                            fullCalcOnLoad=True,forceFullCalc=True,iterate=False)
        self.book.properties.title='VISIT TKY 炭素・窒素・水の行列方程式'
        self.book.properties.subject='旧 visit_local / 日次版 / 数式を持つ監査用ブック'
        self.book.properties.creator='Control_carbon_model'
        self.cache={}
        self.locations={}

    def sheet(self,title,headers=None,note=None):
        w=self.book.create_sheet(title)
        w.sheet_view.showGridLines=False
        w.freeze_panes='D5'
        w.sheet_properties.pageSetUpPr.fitToPage=True
        w.sheet_properties.outlinePr.summaryRight=False
        w.sheet_properties.tabColor=TEAL
        w.cell(1,1,title).font=Font(name='Yu Gothic',size=17,bold=True,color=NAVY)
        w.cell(2,1,note or '数値セルを選択すると、数式バーにExcel関数が表示されます。').font=Font(name='Yu Gothic',size=10,color='555555')
        if title!='00_案内':
            w.cell(1,8,'案内へ').hyperlink="#'00_案内'!A1"
            w.cell(1,8).font=Font(color=BLUE,underline='single')
        if headers:
            for j,h in enumerate(headers,1):
                c=w.cell(4,j,h)
                c.fill=PatternFill('solid',fgColor=NAVY)
                c.font=Font(name='Yu Gothic',bold=True,color='FFFFFF')
                c.alignment=Alignment(wrap_text=True)
        for letter,width in [('A',32),('B',46),('C',27),('D',24),('E',61),('F',72),('G',30)]:
            w.column_dimensions[letter].width=width
        return w

    def formula(self,w,row,column,formula,value,comment=None):
        if len(formula)>8192: raise ValueError(f'Excel formula exceeds 8192: {w.title}:{row},{column}')
        c=w.cell(row,column,formula)
        c.number_format='0.000000;[Red]-0.000000;0'
        c.font=Font(name='Yu Gothic',size=10,color='17252A')
        if comment: c.comment=Comment(comment,'原典と数式')
        self.cache[w.title,c.coordinate]=value
        return c

    def link(self,w,row,column,name):
        return self.formula(w,row,column,f'=IF(ISNUMBER({name}),{name},"{UNKNOWN}")',
                            self.values[name],self.model.graph.nodes[name].source)

    def name(self,key,sheet,address):
        d=DefinedName(key,attr_text=f"'{sheet}'!${re.sub(r'[0-9]','',address)}${re.sub(r'[^0-9]','',address)}")
        self.book.defined_names.add(d)
        self.locations[key]=(sheet,address)

    def table_nodes(self):
        excluded={'11_B','13_A'}
        for title in sorted({n.sheet for n in self.model.graph.nodes.values()}-excluded):
            w=self.sheet(title,['記号（Excel定義名）','説明','単位','採用値／計算値','出典（原典は94_出典）','補足','初期化値：参考のみ'])
            row=5
            for name,node in self.model.graph.nodes.items():
                if node.sheet!=title: continue
                w.cell(row,1,name); w.cell(row,2,node.label); w.cell(row,3,node.unit)
                w.cell(row,5,node.source); w.cell(row,6,node.note)
                if node.expression is None:
                    c=w.cell(row,4,node.value)
                    c.fill=PatternFill('solid',fgColor=AMBER if node.value is None else 'E2EEF9')
                    c.font=Font(color=BLUE)
                    c.number_format='0.000000;[Red]-0.000000;0'
                    c.comment=Comment(node.source+'\n'+node.note,'入力出典')
                else:
                    self.formula(w,row,4,self.model.graph.excel(name),self.values[name],
                                 node.expression+'\n'+node.source+'\n'+node.note)
                for j in (2,3,5,6):w.cell(row,j).alignment=Alignment(wrap_text=True,vertical='top')
                w.row_dimensions[row].height=32
                self.name(name,title,f'D{row}')
                if title=='10_X':
                    state=next(s for s in self.model.states if s.name==name)
                    w.cell(row,7,state.initial)
                    w.cell(row,1).fill=PatternFill('solid',fgColor=GROUP[state.group])
                row+=1
            w.auto_filter.ref=f'A4:G{row-1}'
            if title in ('04_気象入力','05_履歴入力','10_X'):
                dv=DataValidation(type='decimal',operator='between',formula1='-1E+15',formula2='1E+15',allow_blank=True)
                dv.errorTitle='数値を入力してください'; dv.error='未提供なら空欄のままにしてください。'
                dv.showErrorMessage=True; w.add_data_validation(dv); dv.add(f'D5:D{row-1}')

    def matrices(self):
        m=self.model; n=len(m.states)
        for title,key,columns in [('11_B','b',m.mu_names),('13_A','a',[s.name for s in m.states])]:
            w=self.sheet(title,note='列から行への係数。Bのμは正の器官純入力＋沈着等＋数値補正。炭素μ≠GPPそのもの。')
            w.freeze_panes='C5'; w.sheet_view.zoomScale=65
            for j,label in enumerate(columns,3):
                w.cell(4,j,label).alignment=Alignment(textRotation=90)
                w.column_dimensions[col(j)].width=15
            w.row_dimensions[4].height=140
            for i,state in enumerate(m.states):
                w.cell(i+5,1,state.name); w.cell(i+5,2,state.label)
                w.cell(i+5,1).fill=PatternFill('solid',fgColor=GROUP[state.group])
                for j in range(len(columns)):
                    name=f'mx_{key}_{i}_{j}'
                    node=m.graph.nodes[name]
                    cell=self.formula(w,i+5,j+3,m.graph.excel(name),self.values[name],node.expression+'\n'+node.note)
                    if node.expression=='0': cell.font=Font(color='BBC6CA',size=9)
                    elif i==j and key=='a': cell.fill=PatternFill('solid',fgColor=PALE)
                    self.name(name,title,cell.coordinate)
        w=self.sheet('12_mu',['j','入力チャンネル','単位','μ','意味','出典'],
                     'μは異なる単位を持つ入力ベクトル。末尾37成分はX=0等の特異分解残差であり、気象ではない。')
        for j,name in enumerate(m.mu_names,5):
            node=m.graph.nodes[name]
            w.cell(j,1,j-4); w.cell(j,2,name); w.cell(j,3,node.unit)
            self.link(w,j,4,name); w.cell(j,5,node.label); w.cell(j,6,node.source)
        for title,part in [('14_xi','xi'),('15_K','k')]:
            w=self.sheet(title,note=('ξ_j = 総控除 / (K_j X_j)。本ブックのξは環境だけでなく状態・順序・配分にも依存。' if part=='xi'
                                    else 'K：Cは原典基準率、N・水・補助状態は1 day⁻¹の代数的規格化。新しい生物学的パラメータではない。'))
            w.freeze_panes='C5';w.sheet_view.zoomScale=65
            for j,s in enumerate(m.states,3):
                w.cell(4,j,s.name).alignment=Alignment(textRotation=90)
                w.column_dimensions[col(j)].width=15
            w.row_dimensions[4].height=120
            for i,s in enumerate(m.states):
                w.cell(i+5,1,s.name);w.cell(i+5,2,s.label)
                for j in range(n):
                    if i==j: self.link(w,i+5,j+3,'m_'+part+'_'+s.name[2:])
                    else: w.cell(i+5,j+3,0).font=Font(color='BBC6CA')
        w=self.sheet('16_M',note='M=AξK。これは凍結係数の行列であり、非線形系のJacobianではない。容量−M⁻¹Bμを連成平衡と同一視しない。')
        w.freeze_panes='C5'
        for j,s in enumerate(m.states,3):
            w.cell(4,j,s.name).alignment=Alignment(textRotation=90); w.column_dimensions[col(j)].width=15
        w.row_dimensions[4].height=120
        for i,s in enumerate(m.states):
            w.cell(i+5,1,s.name);w.cell(i+5,2,s.label)
            for j,d in enumerate(m.states):
                if m.a_expr[i,j]=='0': w.cell(i+5,j+3,0);continue
                names=[f'mx_a_{i}_{j}','m_xi_'+d.name[2:],'m_k_'+d.name[2:]]
                nums=[self.values[k] for k in names]
                value=nums[0]*nums[1]*nums[2] if all(isinstance(v,(int,float)) for v in nums) else None
                guards=','.join('ISNUMBER('+k+')' for k in names)
                self.formula(w,i+5,j+3,f'=IF(AND({guards}),'+ '*'.join(names)+f',"{UNKNOWN}")',value)

    def integrated(self):
        m=self.model;n=len(m.states)
        w=self.book['01_統合方程式'];w.freeze_panes='D10';w.sheet_view.zoomScale=55
        w.cell(2,1,'dX/dt ≔ fΔ(X,h,z) = B(X,h,z) μ(X,h,z) + A(X,h,z) ξ(X,h,z) K X    [Δ=1 day]')
        w.cell(3,1,'37状態：C 21 + N 11 + 水 3 + 補助 2。hは履歴。これは日次写像の増分表示で、厳密な連続時間埋込みではありません。')
        w.cell(4,1,'左右にスクロールすると全行列が並びます。行名の色：緑=C、紫=N、青=水、橙=補助。単位の異なる行を合計しないでください。')
        w.cell(5,1,'フェノロジーは23、気孔コンダクタンス記憶は22、翌日の履歴対応は93。月・年の更新は日次写像を合成します。')
        w.cell(6,1,'併記する履歴更新： h_next = HΔ(X,h,z)。現在のhへ戻す循環参照は設けません。')
        w.cell(6,7,'履歴更新の値を見る').hyperlink="#'93_計算順序'!A15"
        w.cell(7,1,'状態');w.cell(7,2,'説明');w.cell(7,3,'dX/dt')
        for i,s in enumerate(m.states,10):
            w.cell(i,1,s.name);w.cell(i,2,s.label);self.link(w,i,3,'mx_rhs_'+s.name[2:])
            w.cell(i,1).fill=PatternFill('solid',fgColor=GROUP[s.group])
        cursor=5
        blocks=[('B','11_B',n,len(m.mu_names)),('μ','12_mu',len(m.mu_names),1),
                ('A','13_A',n,n),('ξ','14_xi',n,n),('K','15_K',n,n),('X','10_X',n,1)]
        for title,source,rows,cols in blocks:
            w.cell(7,cursor,title).font=Font(size=20,bold=True,color=TEAL)
            for j in range(cols):
                w.column_dimensions[col(cursor+j)].width=14
                label=(m.mu_names[j] if title=='B' else m.states[j].name if title in ('A','ξ','K') else title)
                w.cell(9,cursor+j,label).alignment=Alignment(textRotation=90)
            for i in range(rows):
                for j in range(cols):
                    addr=f'{col(j+3)}{i+5}' if cols>1 else f'D{i+5}'
                    source_cell=self.book[source][addr]
                    value=self.cache.get((source,addr),source_cell.value)
                    target=f"'{source}'!{addr}"
                    self.formula(w,i+10,cursor+j,f'=IF(ISNUMBER({target}),{target},"{UNKNOWN}")',value)
            w.cell(7,cursor-1,'+' if title=='A' else '=' if title=='B' else '×').font=Font(size=20,bold=True)
            cursor+=cols+2
        w.row_dimensions[9].height=150
        w.print_options.horizontalCentered=False
        w.print_area=f'A1:{col(cursor)}{max(n,len(m.mu_names))+10}'

    def catalog(self):
        m=self.model
        w=self.sheet('40_フラックス',['フラックス記号','説明','単位','値','供与元 → 受取係数','出典','注記'],
                     '全132チャンネル。供与元なしはBμ、供与元ありはAξKX。ただしX=0等の残差はBμへ分離。')
        for row,f in enumerate(m.fluxes,5):
            node=m.graph.nodes[f.name]
            for j,value in [(1,f.name),(2,node.label),(3,node.unit),(5,(f.donor or '系外・有効入力')+' → '+str(f.recipients)),(6,node.source),(7,f.note)]:w.cell(row,j,value)
            self.link(w,row,4,f.name)
        w.auto_filter.ref=f'A4:G{len(m.fluxes)+4}'
        for title,group in [('41_炭素','C'),('42_窒素','N'),('43_水','W')]:
            w=self.sheet(title,['状態','説明','単位','現在X','Bμ','AξKX','dX/dt','X_next','行列−台帳'],
                         '炭素はC4仮想状態を含む。生態系炭素合計ではC4重み0。Nは原典数値規約であり単位疑義を99に明記。')
            row=5
            for s in m.states:
                if s.group!=group:continue
                w.cell(row,1,s.name);w.cell(row,2,s.label);w.cell(row,3,s.unit)
                for j,key in [(4,s.name),(5,'mx_input_'+s.name[2:]),(6,'mx_transfer_'+s.name[2:]),(7,'mx_rhs_'+s.name[2:]),(8,'m_end_'+s.name[2:]),(9,'mx_error_'+s.name[2:])]:self.link(w,row,j,key)
                row+=1
        w=self.sheet('90_変数辞書',['記号','説明','単位','現在値','出典','短い記号式／注記','数式セルへ'])
        for row,(name,node) in enumerate(m.graph.nodes.items(),5):
            for j,value in [(1,name),(2,node.label),(3,node.unit),(5,node.source),(6,node.expression or node.note)]:w.cell(row,j,value)
            self.link(w,row,4,name)
            title,addr=self.locations[name]
            w.cell(row,7,f'{title}!{addr}').hyperlink=f"#'{title}'!{addr}"
        w.auto_filter.ref=f'A4:G{len(m.graph.nodes)+4}'
        w=self.sheet('93_計算順序',['順番','処理／履歴','入力','出力','注意'])
        stages=[('1','パラメータ・気象・状態・前日記憶','02,03,04,05,10','当日の評価条件','未設定を0扱いしない'),
                ('2','SLA、LAI、放射・気圧・PM','状態・前日フェノロジー','20','日平均/正午を原典どおり区別'),
                ('3','雪・遮断・2層水収支','前日のgc','21','下流の当日gcを参照しない'),
                ('4','単葉光合成と気孔の6回更新','更新済み水量','22','6段階を展開。循環参照なし'),
                ('5','N獲得、フェノロジー、植物C、配分N','GDD/CDDなど','23,24,26','同日途中の状態を専用セルで保持'),
                ('6','土壌C→無機化→DOC/CASA→ガス→N収支','植物リター、更新済み土壌C','25,26,27','N無機化分母は更新済みC'),
                ('7','フラックス台帳→行列要素→行列積','全132フラックス','11–17,30,40','同じXを入力。更新先をXへ自動参照させない'),
                ('8','次の日への受け渡し','m_end_* / q_* / pht_*_gc','次の日のX/h','このブックは単一日を評価。履歴は別日の値として渡す')]
        for r,row in enumerate(stages,5):
            for j,value in enumerate(row,1):w.cell(r,j,value)
        r=15
        for p in ('t','g','v'):
            for key in ('gdd','cdd','flush','shed','gc'):
                new=f'pht_{p}_gc' if key=='gc' else f'q_{p}_{key}'
                w.cell(r,1,'翌日');w.cell(r,2,f'h_{p}_{key}');w.cell(r,3,new)
                self.link(w,r,4,new);w.cell(r,5,'左の値を次の日の履歴へ渡す（同日の参照は禁止）');r+=1

    def documentation(self,verification=None):
        m=self.model
        w=self.book['00_案内']
        entries=[
            ('目的','TKY落葉広葉樹林について、日次VISITのC・N・水を一つの行列式からたどる監査用xlsx。'),
            ('最初に開く','01_統合方程式。次に41_炭素 / 42_窒素 / 43_水で部分系を見る。'),
            ('最重要：未入力','Configは気象・現在状態・履歴を与えない。該当セルは空欄、依存値は「未設定/定義域」。推測値は入れていない。'),
            ('入力する場所','04_気象入力 / 05_履歴入力 / 10_X の黄色D列。02はConfig値、03は原典定数。青字は入力値。'),
            ('数式を見る','値セルを選択し、Excel数式バーを見る。コメントには短い記号式と原典。90_変数辞書から相互移動できる。'),
            ('行列の意味','B：有効入力先、μ：正の器官純入力等、A：内部輸送と控除、ξ：状態・環境依存倍率、K：基準日率、X：現在状態。'),
            ('μについて','維持呼吸を含むEPPから正の器官純入力を抽出。GPPそのものではない。ゼロ丸めと特異分解残差も別チャンネルに含む。'),
            ('分解は一意ではない','Kの取り方でξが変わる。本ブックのN・水K=1/dayは代数的規格化であって寿命の仮定ではない。'),
            ('時間の意味','fΔ=(FΔ−X)/Δ, Δ=1日。X_next=X+ΔfΔは元の段階計算の増分。連続ODEの厳密な埋込みや月次モデルとは呼ばない。'),
            ('月次・年次','月・年は日次写像を逐次合成する。Δを30や365に書き換えない。平均気象の代入は非線形性により平均速度とは一般に異なる。'),
            ('履歴と閉じた系','37状態だけでは閉じない。GDD/CDD、展葉・落葉日数、前日gcをhとして横に置く。93に翌日の対応を明記。'),
            ('Luo形式との違い','外形はBμ+AξKX。係数はXに依存し、元モデルの順序依存・補正も含む。M=AξKはJacobianではない。'),
            ('capacityとQSE','−M(X,z)⁻¹B(X,z)μ(X,z)は凍結診断。連成QSEはf(X,h,z)=0等を解く必要があり、同一とは限らない。'),
            ('質量単位','CはMgC/ha、Nは原典収支のgN/ha、水はmm。DOC濃度/CASA記憶は補助状態で、物理総量へ足さない。'),
            ('選択範囲','旧visit_local・TKY・DBF・日次式FLUX_SCHEME=0・N_CYCLE=1。他サイト/30分版/農業/湿地/同位体は非選択。'),
            ('草本','TKYのC3林床は有効。C4面積は0だが原典のN更新が残るため仮想群を表示。任意には削除していない。'),
            ('原典の疑義','baseflow二重控除、N単位混在などを保持し99へ記録。「原典どおり」と「科学的に正しい」を区別する。'),
            ('ゼロ状態','X_j=0でも同日流入後の流出があり得る。ξ=F/(KX)の0割りを避け、残差をBμに明示する。'),
            ('再計算','標準Excel数式と定義名を使用し自動計算ON、反復計算OFF。保存済み計算値も同梱。Excelアプリ実機検証は未実施。'),
            ('検証と限界','98_検証を参照。選択経路を数式化した研究用監査版。全分岐・全パラメータ領域の完全な原典同値を保証しない。'),
            ('データ境界',f'Sachitama2001/VISIT-matrix @ {SOURCE_COMMIT}; Config SHA256={m.config_hash}'),
        ]
        if hasattr(m, 'baseline'):
            b=m.baseline
            entries[2]=('実データ基準日',f"{b['date']}。未入力{b['before']}→{b['after']}。06に換算・出典、07に未採用restart、08にファイル照合。現在Xは未提供。")
        for r,(a,b) in enumerate(entries,5):
            w.cell(r,1,a);w.cell(r,2,b);w.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top');w.row_dimensions[r].height=34
        w.column_dimensions['B'].width=118
        r=29
        for title in self.book.sheetnames:
            w.cell(r,1,title).hyperlink=f"#'{title}'!A1";r+=1
        w=self.sheet('92_設定と省略',['設定／範囲','採用','理由','出典／注意'])
        rows=[('site / veg_type','TKY / 4','指定どおり','Config Site C4,C5'),('FLUX_SCHEME','0 日次','最も簡単な実験経路','元setting.hの1から意図的に切替。30分FvCB/熱収支は含めない'),
              ('N_CYCLE','1','C–N相互作用を含める','definition.h'),('AG_LMA','1','TKYの季節SLAを維持','leaf_aging.c'),
              ('SCI_SCHEME / WH_CH4','0 / 0','同位体/湿地を非選択','TKY森林・簡単設定'),('NECB_DOC/POC/METHANE/BVOC','0','診断排出とC収支を混同しない','DOC/VOC/CH4式は診断シートに残す'),
              ('FIX_PHENOLOGY / F_DOWNREG / EX_*','0','固定フェノロジー/CO2増倍/加熱/乾燥実験なし','setting.h'),
              ('C3/C4面積','1 / 0','TKYサイト分岐','C4仮想状態のN処理は99参照'),
              ('crop・harvest・施肥','非選択','森林設定','phenology_crop, fertilizationほか'),
              ('spinup/read/write/alloc/clear','物理過程としては省略','計算機処理','clearによる0初期化は有効式の意味に反映'),
                ('init値','参考列のみ','spinup種値≠現在の森林量','initi_mass.c; native診断指定時は10_Xに基準日開始値'),
              ('Hikosaka/Anten・game','非選択','呼出しがコメント化','location_proc.c'),
              ('大気CO2/CH4・N沈着','外部の日値を入力','実データを捏造しない','読込方法や年補間を独立した生物過程にはしない'),
              ('未知値','99へ列挙','代入しない','空欄と0を区別する')]
        for r,row in enumerate(rows,5):
            for j,val in enumerate(row,1):w.cell(r,j,val)
        w=self.sheet('94_出典',['ファイル','SHA256','境界','備考'])
        for r,(filename,txt) in enumerate(sorted(m.source_text.items()),5):
            w.cell(r,1,'visit_local/'+filename);w.cell(r,2,hashlib.sha256((m.source/filename).read_bytes()).hexdigest());w.cell(r,3,SOURCE_COMMIT)
        r=w.max_row+1;w.cell(r,1,str(m.config));w.cell(r,2,m.config_hash);w.cell(r,3,'指定Config実ファイル')
        for title,original in [('95_Config_Site','Site'),('96_Config_Daily','Parameter-daily'),('97_Config_30min','Parameter-30min')]:
            w=self.sheet(title,note='原データの値スナップショット（元の行番号+4）。30minは非選択・参照用。採用値の編集は02_パラメータで行う。')
            for row in m.config_book[original]:
                for c in row:
                    if c.value is not None:w.cell(c.row+4,c.column,c.value)
        w=self.sheet('98_検証',['項目','結果','方法／限界'])
        checks=[('循環参照','なし（DAG）','全記号式は登録済み上流ノードだけを参照。X_next→Xのリンクなし。'),
                ('Excel式長','8192文字以下','生成時に全数検査。標準関数のみ。'),
                ('保存済み計算値','全数式セルへ格納','openpyxl再読込data_onlyと数式読込を別々に検査。'),
                ('欠測値','空欄＋文字表示','推測値0で計算を継続しない。IF枝に依存しない欠測は伝播しない。'),
                ('実Excel再計算','未実施','この環境にExcel/LibreOfficeなし。Python式木評価とOOXML構造検査を実施。'),
                ('原典全分岐同値','未保証','入力定義域、未初期化変数、浮動小数差、無効設定の範囲外を含む。'),
                ('状態数 / フラックス数',f'{len(m.states)} / {len(m.fluxes)}','C21,N11,W3,補助2。'),
                ('状態値の根拠','native開始値' if getattr(m,'baseline',{}).get('native_rows') else '未提供',
                 'native診断指定時は2000-07-15の日次更新前。人工ケースは成果xlsxへ代入しない。')]
        if getattr(m,'baseline',{}).get('native_rows'):
            checks.append(('native 1日照合最大規格化誤差',f"{m.baseline['native_max_scaled_error']:.3e}",
                           '37状態。max(|Excel終了−C終了|/max(1,|C終了|))。98A参照。'))
        for k,v in (verification or {}).items():checks.append((k,str(v),'自動テスト結果。再現方法はdocs/visit_matrix_workbook.md'))
        for r,row in enumerate(checks,5):
            for j,val in enumerate(row,1):w.cell(r,j,val)

    def process_inventory(self):
        m=self.model
        w=self.sheet('91_過程網羅',['ファイル:行','関数','分類','数式セル数','扱い／理由'])
        nonselected={'canopy_proc.c','leaf_energy.c','photosynthesis_30min.c','canopy_scheme.c','game_optimization.c',
                     'canopy.c','two_canopy.c','isotope.c','agriculture.c','disturbance.c','parameter_var.c'}
        bundled={
            'f_mortality':'lf0/lc0/lr0を当日率へ戻す。c_*_lf_rateおよびc_*_lc/lrへ展開。',
            'f_lf':'通常落葉率×展葉後葉量。c_*_lfへ展開。',
            'f_solhgt':'e_sinnoonおよび48個のe_swへ展開。',
            'f_toprad':'e_dtcおよび48個のe_swへ展開。',
            'f_ppfd':'Black/Tooming/McCree/Dyeの計算を20へ展開。',
            'f_resp_h':'非選択の呼吸補助関数。採用経路はf_rfm/f_rcm/f_rrmとf_rfg/f_rcg/f_rrg。',
            'f_leaf_aging':'Vcmax/Jmax/暗呼吸の季節変化は30分経路向け。日次に影響するSLAはf_sla_changeで採用。',
            'f_n_deposit':'呼出しがコメント化。現用の日沈着値z_depo_no3/nh4を外部入力。',
            'co2_in_canopy':'overs/undersの30分生理用CO2。日次f_ecophysiologyはloct.aCO2を使用。',
            'atmco2_trend':'外部大気シナリオの年次読込・補間。評価日のCO2/CH4を明示入力。',
            'f_experiment':'時間ループ・実験制御。自然森林、攪乱なし、感度摂動なし。',
            'f_flux_site':'サイト駆動・出力ループ。生物過程はdaily_scheme側で表現。',
            'f_ansis_mon':'月次統計集計。状態方程式の新たな過程ではない。',
            'f_ansis_ann':'年次統計集計。状態方程式の新たな過程ではない。',
            'f_doyTmody':'暦変換。0始まりDOYを外部入力。',
            'f_modyTdoy':'暦変換。0始まりDOYを外部入力。',
            'delete_plant':'メモリ/状態初期化。生物学的枯死とは異なる。',
            'delete_soil':'メモリ/状態初期化。分解とは異なる。',
            'pflux_zero':'毎日フラックス変数を0へ戻す。計算順の意味として採用。',
            'sflux_zero':'毎日フラックス変数を0へ戻す。計算順の意味として採用。',
            'set_parameter':'02_パラメータへConfigの値を直接読込。その他PFTは非選択。',
            'f_setting':'実験設定の読込。92_設定と省略に選択値を固定。',
        }
        function_refs={}
        for node in m.graph.nodes.values():
            if '::' in node.source:
                filename=node.source.split('/')[1].split(':')[0];fn=node.source.split('::')[-1]
                function_refs[filename,fn]=function_refs.get((filename,fn),0)+1
        admin=('init','read','write','clear','spin','alloc_mem','output','input','open','close','free','print','save','main')
        r=5
        for filename,source in sorted(m.source_text.items()):
            if not filename.endswith('.c'):continue
            for match in re.finditer(r'^[ \t]*(?:void|double|int|long|short)\s+(\w+)\s*\(',source,re.M):
                fn=match.group(1);num=function_refs.get((filename,fn),0)
                if num:status,why='式を展開','対応する数式は90_変数辞書で出典を検索'
                elif fn in bundled:status,why='集約／明示的非選択',bundled[fn]
                elif fn.startswith('f_clim_correct_'):
                    status,why='外部気象の前処理','04の気象は補正後の値として定義。TKY以外の補正は非選択。補正係数は生物過程へ二重に入れない。'
                elif any(t in fn.lower() for t in admin):status,why='入出力・初期化等','生物過程として独立した行列要素を作らない。参照定数や実行順は別記。'
                elif filename in nonselected or any(t in fn.lower() for t in ('fvbc','fvbc','farquhar','newton','isotope','13c','walter','crop','bareland','hikosaka','fertil','game')):
                    status,why='非選択経路','日次・TKY・DBF・SCI=0・WH_CH4=0等の選択により不採用'
                elif fn in ('daily_scheme','f_loct_proc','plant_process','f_cycle_soil','f_ecophysiology'):
                    status,why='処理順・集約','段階計算とフラックス台帳へ分解'
                else:status,why='補助／分岐監査対象','独立セル未対応。別関数での式展開・診断集約・非選択経路を含む。全関数転記済みとは主張しない。'
                w.cell(r,1,f'visit_local/{filename}:{source.count(chr(10),0,match.start())+1}')
                w.cell(r,2,fn);w.cell(r,3,status);w.cell(r,4,num);w.cell(r,5,why);r+=1
        w.auto_filter.ref=f'A4:E{r-1}'

    def unknowns(self):
        m=self.model
        w=self.sheet('99_不明と要確認',['ID／記号','項目','問題／単位','必要な情報','参照先'],
                     '最後のシート。未提供値と原典の疑義を区別し、未知のデフォルト値は埋めない。')
        r=5
        for code,subject,reason,needed in m.issues:
            for j,val in enumerate((code,subject,reason,needed,''),1):w.cell(r,j,val)
            r+=1
        for name,node in m.graph.nodes.items():
            if node.expression is not None or node.value is not None:continue
            title,addr=self.locations[name]
            for j,val in enumerate((name,node.label,node.unit,'評価日の値を提供（未定義分岐は修正ソースが必要）',f'{title}!{addr}'),1):w.cell(r,j,val)
            w.cell(r,5).hyperlink=f"#'{title}'!{addr}";r+=1
        w.auto_filter.ref=f'A4:E{r-1}'

    def build(self,path,verification=None):
        self.sheet('00_案内',['項目','説明'])
        self.sheet('01_統合方程式')
        self.table_nodes();self.matrices();self.integrated();self.catalog()
        self.process_inventory()
        if hasattr(self.model, 'baseline'):
            b=self.model.baseline
            for title, rows, headers, note in [
                ('06_基準日データ', b['imports']+b['references'],
                 ['記号／参考量','値','単位','データ出典','換算・採否'],
                 b['date']+'。採用入力は固定値スナップショット。下段の参考量は状態に代入しない。式・Configは旧版を維持。'),
                ('07_restart参照', b['restart'],
                 ['対応状態','参考値','単位','データ出典','採否'],
                 'spinup終了時の21量。夏の貯留と混同しないため10_Xには代入していない。'),
                ('08_データ来歴', b['manifest'],
                 ['実ファイル','SHA256','旧visit_localとのファイル比較'],
                 'データは読取専用。ファイル一致は記録するが、不一致をモデル同値とみなさない。')]:
                w=self.sheet(title,headers,note)
                for r,row in enumerate(rows,5):
                    for j,value in enumerate(row,1):
                        w.cell(r,j,value)
                        w.cell(r,j).alignment=Alignment(wrap_text=True,vertical='top')
                w.column_dimensions['A'].width=45
                w.column_dimensions['D'].width=65
                w.column_dimensions['E'].width=95
            if b.get('native_rows'):
                w=self.sheet('09_native実行',['記号','当日開始','f_loct_proc後','日次更新後','単位','採用規則'],
                             '2000-07-15の3時点。X/hは開始、気象zはf_loct_proc後を入力へ採用。')
                for r,row in enumerate(b['native_rows'],5):
                    for j,value in enumerate(row,1):w.cell(r,j,value)
                w=self.sheet('98A_native照合',['状態','native開始','native終了','Excel式終了','差','規格化誤差','単位'],
                             '37状態の1日更新照合。保水容量はSaxton式後の実行値。')
                for r,row in enumerate(b['native_validation'],5):
                    for j,value in enumerate(row,1):w.cell(r,j,value)
        self.documentation(verification);self.unknowns()
        self.book._sheets.sort(key=lambda w:w.title)
        for r,title in enumerate(self.book.sheetnames,29):
            self.book['00_案内'].cell(r,1,title).hyperlink=f"#'{title}'!A1"
        for w in self.book:
            w.sheet_properties.pageSetUpPr.fitToPage=True
        path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
        save_with_caches(self.book,path,self.cache)
        return audit_workbook(path,self)


def audit_workbook(path,exporter=None):
    formula=load_workbook(path,data_only=False)
    data=load_workbook(path,data_only=True)
    count=0;maxlen=0
    for w in formula:
        for row in w:
            for c in row:
                if c.data_type!='f':continue
                assert '#REF!' not in c.value,(w.title,c.coordinate)
                assert data[w.title][c.coordinate].value is not None,(w.title,c.coordinate)
                maxlen=max(maxlen,len(c.value));count+=1
    assert maxlen<=8192
    assert formula.calculation.iterate is False
    assert not formula._external_links
    assert formula.sheetnames[-1]=='99_不明と要確認'
    if exporter:
        for name,n in exporter.model.graph.nodes.items():
            title,addr=exporter.locations[name]
            if n.expression is None:
                saved=data[title][addr].value
                assert (math.isclose(saved,n.value,rel_tol=1e-14,abs_tol=1e-14)
                        if isinstance(saved,(int,float)) and isinstance(n.value,(int,float))
                        else saved==n.value),(name,saved,n.value)
    return {'sheets':len(formula.sheetnames),'formula_cells':count,'max_formula_length':maxlen,
            'defined_names':len(formula.defined_names),'file_bytes':Path(path).stat().st_size}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=SOURCE_DEFAULT)
    parser.add_argument('--config',type=Path)
    parser.add_argument('--output',type=Path,default=Path('artifacts/visit_matrix_workbook/VISIT_TKY_matrix_equations.xlsx'))
    parser.add_argument('--verification',type=Path)
    parser.add_argument('--baseline-data',type=Path,help='visitb実データのINPUT/OUTPUTを含むディレクトリ')
    parser.add_argument('--baseline-date',default='2000-07-15')
    parser.add_argument('--native-snapshot',type=Path,help='ansis診断版の3時点TSV')
    args=parser.parse_args()
    verification=json.loads(args.verification.read_text()) if args.verification else None
    model=TKYModel(args.source,args.config)
    if args.baseline_data:
        from .visit_workbook_data import load_baseline
        load_baseline(model,args.baseline_data,args.baseline_date,args.native_snapshot)
    exporter=Exporter(model)
    print(json.dumps(exporter.build(args.output,verification),ensure_ascii=False,indent=2))


if __name__=='__main__': main()
